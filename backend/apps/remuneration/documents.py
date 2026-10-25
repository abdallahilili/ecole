"""Fiche de paiement (PDF), etat de paiement mensuel (PDF) et exports Excel."""

from collections import defaultdict
from decimal import Decimal

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib.units import cm
from reportlab.platypus import Spacer

from apps.commun.excel import orienter_feuille
from apps.commun.traduction import traduire
from apps.commun.pdf import Tableau as Table, STYLE_PETIT, STYLE_TITRE, construire_pdf, entete_etablissement, gras, p, style_tableau
from . import services
from .models import StatutPaiement, StatutReleve, TypeAjustement

LIBELLES_PAIEMENT = {StatutPaiement.NON_PAYE: "Non paye", StatutPaiement.PARTIEL: "Partiel", StatutPaiement.PAYE: "Paye"}
MOIS = ["janvier", "fevrier", "mars", "avril", "mai", "juin", "juillet", "aout", "septembre", "octobre", "novembre", "decembre"]


def _libelle_paiement(statut, valide):
    return traduire(LIBELLES_PAIEMENT[statut]) + ("" if valide else " " + traduire("(provisoire)"))


def um(valeur):
    """1500 -> '1 500,00 UM' (convention francaise)."""
    texte = f"{Decimal(valeur):,.2f}".replace(",", " ").replace(".", ",")
    return f"{texte} {traduire('UM')}"


def nombre(valeur):
    return f"{Decimal(valeur):.2f}".replace(".", ",")


def nom_mois(mois):
    nom = traduire(MOIS[mois.month - 1])
    return f"{nom.capitalize()} {mois.year}"


def pdf_fiche(releve):
    t = services.totaux(releve)
    professeur = releve.professeur
    lignes = list(releve.lignes.select_related("classe", "matiere"))
    horaire = releve.professeur.mode_remuneration == "HORAIRE"
    elements = entete_etablissement() + [p(f"FICHE DE PAIEMENT - {nom_mois(releve.mois)}", STYLE_TITRE)]
    mode = "Au nombre d'heures" if horaire else "Salaire fixe"
    statut = "Valide" if releve.statut == StatutReleve.VALIDE else "Provisoire (non valide)"
    infos = Table([
        [p(f"Professeur : {professeur.nom_complet}"), p(f"Matricule : {professeur.matricule}")],
        [p(f"Mode de remuneration : {mode}"), p(f"Statut du releve : {statut}")],
    ], colWidths=[9 * cm, 9 * cm])
    elements += [infos, Spacer(1, 8)]

    donnees = [["Date", "Classe", "Matiere", "Heures", "Tarif", "Montant"]]
    for l in lignes:
        donnees.append([f"{l.date:%d/%m/%Y}", p(l.classe.nom), p(l.matiere.nom_affiche), nombre(l.heures),
                        nombre(l.tarif) if horaire else "-", nombre(l.montant) if horaire else "-"])
    if len(donnees) == 1:
        donnees.append(["", p("Aucune seance"), "", "", "", ""])
    tableau = Table(donnees, repeatRows=1, colWidths=[2.4 * cm, 3.6 * cm, 5 * cm, 2 * cm, 2.4 * cm, 2.6 * cm])
    tableau.setStyle(style_tableau())
    elements += [tableau, Spacer(1, 8)]

    recap = defaultdict(lambda: Decimal("0"))
    for l in lignes:
        recap[(l.matiere.nom_affiche, l.classe.nom)] += l.heures
    if recap:
        lignes_recap = [["Matiere", "Classe", "Heures"]] + [
            [p(m), p(c), nombre(h)] for (m, c), h in sorted(recap.items())]
        table_recap = Table(lignes_recap, colWidths=[7 * cm, 5 * cm, 3 * cm], repeatRows=1)
        table_recap.setStyle(style_tableau())
        elements += [p("Recapitulatif par matiere et par classe", STYLE_PETIT), table_recap, Spacer(1, 8)]

    bilan = [[gras("Heures effectuees"), nombre(t["heures"])], [gras("Montant brut"), um(t["brut"])]]
    for a in releve.ajustements.all():
        signe = "+" if a.type == TypeAjustement.PRIME else "-"
        bilan.append([p(traduire("Prime" if a.type == TypeAjustement.PRIME else "Retenue") + f" : {a.libelle}"), f"{signe} {um(a.montant)}"])
    bilan += [[gras("Montant net a payer"), gras(um(t["net"]))]]
    table_bilan = Table(bilan, colWidths=[11 * cm, 5 * cm])
    table_bilan.setStyle(style_tableau(entete=False))
    elements += [table_bilan, Spacer(1, 8)]

    paiements = list(releve.paiements.order_by("date_paiement", "id"))
    if paiements:
        lignes_p = [["Reference", "Date", "Type", "Montant"]] + [
            [pa.reference, f"{pa.date_paiement:%d/%m/%Y}", pa.get_type_display(), um(pa.montant)] for pa in paiements]
        table_p = Table(lignes_p, colWidths=[4 * cm, 3 * cm, 5 * cm, 4 * cm], repeatRows=1)
        table_p.setStyle(style_tableau())
        elements += [p("Paiements", STYLE_PETIT), table_p, Spacer(1, 6)]
    elements.append(gras(f"Total paye : {um(t['paye'])} - Solde restant : {um(t['solde'])}"))
    signatures = Table([[p("Le professeur"), p("Le comptable"), p("Le directeur")]], colWidths=[6 * cm] * 3, rowHeights=[2 * cm])
    elements += [Spacer(1, 14), signatures]
    return construire_pdf(elements, titre="Fiche de paiement",
                          filigrane=None if releve.statut == StatutReleve.VALIDE else "PROVISOIRE")


def _lignes_etat(releves):
    resultats = []
    for r in releves:
        t = services.totaux(r)
        resultats.append((r, t))
    return resultats


def pdf_etat(releves, mois):
    en_tete = ["Professeur", "Heures", "Brut", "Primes", "Retenues", "Net", "Paye", "Solde", "Paiement"]
    donnees = [en_tete]
    total = defaultdict(lambda: Decimal("0"))
    for r, t in _lignes_etat(releves):
        donnees.append([p(r.professeur.nom_complet), nombre(t["heures"]), nombre(t["brut"]), nombre(t["primes"]), nombre(t["retenues"]),
                        nombre(t["net"]), nombre(t["paye"]), nombre(t["solde"]),
                        _libelle_paiement(t["statut_paiement"], r.statut == StatutReleve.VALIDE)])
        for cle in ("heures", "brut", "primes", "retenues", "net", "paye", "solde"):
            total[cle] += t[cle]
    donnees.append([gras("TOTAL")] + [gras(nombre(total[c])) for c in ("heures", "brut", "primes", "retenues", "net", "paye", "solde")] + [""])
    tableau = Table(donnees, repeatRows=1, colWidths=[6 * cm, 2.2 * cm, 2.8 * cm, 2.5 * cm, 2.5 * cm, 2.8 * cm, 2.8 * cm, 2.8 * cm, 3.6 * cm])
    tableau.setStyle(style_tableau())
    elements = entete_etablissement() + [p(f"ETAT DE PAIEMENT DES PROFESSEURS - {nom_mois(mois)}", STYLE_TITRE), tableau,
                                         Spacer(1, 6), p("Montants en UM. Seuls les releves valides font foi.", STYLE_PETIT)]
    return construire_pdf(elements, paysage=True, titre="Etat de paiement")


def xlsx_etat(releves, mois):
    classeur = openpyxl.Workbook()
    feuille = classeur.active
    feuille.title = traduire("Etat de paiement")
    orienter_feuille(feuille)
    feuille.append([traduire(f"Etat de paiement des professeurs - {nom_mois(mois)}")])
    feuille["A1"].font = Font(bold=True, size=12)
    feuille.append([])
    entete = [traduire(h) for h in ["Professeur", "Matricule", "Heures", "Brut", "Primes", "Retenues", "Net", "Paye", "Solde", "Statut du paiement", "Statut du releve"]]
    feuille.append(entete)
    for index in range(1, len(entete) + 1):
        cellule = feuille.cell(row=3, column=index)
        cellule.font = Font(bold=True)
        cellule.fill = PatternFill("solid", fgColor="EEF1F3")
    total = defaultdict(lambda: Decimal("0"))
    detail = []
    for r, t in _lignes_etat(releves):
        feuille.append([r.professeur.nom_complet, r.professeur.matricule] + [float(t[c]) for c in ("heures", "brut", "primes", "retenues", "net", "paye", "solde")]
                       + [traduire(LIBELLES_PAIEMENT[t["statut_paiement"]]), traduire("Valide" if r.statut == StatutReleve.VALIDE else "Brouillon")])
        for cle in ("heures", "brut", "primes", "retenues", "net", "paye", "solde"):
            total[cle] += t[cle]
        for l in r.lignes.select_related("classe", "matiere"):
            detail.append([r.professeur.nom_complet, l.date.isoformat(), l.classe.nom, l.matiere.nom_affiche, float(l.heures), float(l.tarif), float(l.montant)])
    feuille.append([traduire("TOTAL"), ""] + [float(total[c]) for c in ("heures", "brut", "primes", "retenues", "net", "paye", "solde")])
    for cellule in feuille[feuille.max_row]:
        cellule.font = Font(bold=True)
    for index, largeur in enumerate([30, 12, 10, 14, 12, 12, 14, 14, 14, 20, 16], start=1):
        feuille.column_dimensions[get_column_letter(index)].width = largeur

    feuille_detail = classeur.create_sheet(traduire("Detail des seances"))
    orienter_feuille(feuille_detail)
    feuille_detail.append([traduire(h) for h in ["Professeur", "Date", "Classe", "Matiere", "Heures", "Tarif", "Montant"]])
    for cellule in feuille_detail[1]:
        cellule.font = Font(bold=True)
        cellule.fill = PatternFill("solid", fgColor="EEF1F3")
    for ligne in detail:
        feuille_detail.append(ligne)
    for index, largeur in enumerate([30, 12, 14, 24, 10, 12, 14], start=1):
        feuille_detail.column_dimensions[get_column_letter(index)].width = largeur
    return classeur
