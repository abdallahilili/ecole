"""Impression et export des emplois du temps (PDF paysage, Excel)."""

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.units import cm
from reportlab.platypus import PageBreak, Spacer, Table

from apps.commun.pdf import STYLE_PETIT, STYLE_TITRE, construire_pdf, entete_etablissement, multiligne, p, style_tableau
from . import services

MODES = ("classe", "professeur", "salle", "matiere")


def _lignes_seance(s, mode):
    classe = s["classe_nom"] + (f" ({s['groupe_nom']})" if s["groupe_nom"] else "")
    salle = f"Salle {s['salle_nom']}" if s["salle_nom"] else "Sans salle"
    if mode == "classe":
        return [s["matiere_nom"], s["professeur_nom"], salle]
    if mode == "professeur":
        return [s["matiere_nom"], classe, salle]
    if mode == "salle":
        return [s["matiere_nom"], classe, s["professeur_nom"]]
    if mode == "matiere":
        return [classe, s["professeur_nom"], salle]
    return [s["matiere_nom"], classe, s["professeur_nom"]]


def _titre(emploi, mode, nom):
    cible = {"classe": "Classe", "professeur": "Professeur", "salle": "Salle", "matiere": "Matiere"}.get(mode)
    return f"Emploi du temps - {cible} {nom}" if cible else f"Emploi du temps - {emploi.nom}"


def _entites(emploi, mode):
    """[(id, nom)] des entites presentes dans l'emploi, pour l'impression 'toutes les classes'."""
    donnees = services.vue_emploi(emploi, {})
    champ, nom = {"classe": ("classe", "classe_nom"), "professeur": ("professeur", "professeur_nom"),
                  "salle": ("salle", "salle_nom"), "matiere": ("matiere", "matiere_nom")}[mode]
    vus = {}
    for s in donnees["seances"]:
        if s[champ]:
            vus[s[champ]] = s[nom]
    return sorted(vus.items(), key=lambda t: t[1])


def _pages(emploi, filtres, toutes):
    """Liste de (titre, mode, donnees de grille) a imprimer."""
    if toutes in MODES:
        return [(_titre(emploi, toutes, nom), toutes, services.vue_emploi(emploi, {toutes: identifiant}))
                for identifiant, nom in _entites(emploi, toutes)]
    mode = next((m for m in MODES if filtres.get(m)), None)
    donnees = services.vue_emploi(emploi, filtres)
    nom = ""
    if mode:
        champ = {"classe": "classe_nom", "professeur": "professeur_nom", "salle": "salle_nom", "matiere": "matiere_nom"}[mode]
        nom = donnees["seances"][0][champ] if donnees["seances"] else ""
    return [(_titre(emploi, mode, nom), mode, donnees)]


def _index_creneaux(donnees):
    """(jour, index_creneau_debut) -> [seances] et nombre de creneaux couverts."""
    cadre = donnees["cadre"]
    debut_h, debut_m = map(int, cadre["debut"].split(":"))
    origine = debut_h * 60 + debut_m
    par_depart = {}
    for s in donnees["seances"]:
        h, m = map(int, s["heure_debut"].split(":"))
        h2, m2 = map(int, s["heure_fin"].split(":"))
        depart = (h * 60 + m - origine) // cadre["duree_creneau"]
        longueur = max(1, (h2 * 60 + m2 - h * 60 - m) // cadre["duree_creneau"])
        par_depart.setdefault((s["jour"], depart), []).append((s, longueur))
    return par_depart


def pdf_emploi(emploi, filtres, toutes=None):
    elements = []
    for numero, (titre, mode, donnees) in enumerate(_pages(emploi, filtres, toutes)):
        if numero:
            elements.append(PageBreak())
        cadre = donnees["cadre"]
        jours = cadre["jours"]
        par_depart = _index_creneaux(donnees)
        entete = ["Horaire"] + [j["nom"] for j in jours]
        rows, styles = [entete], []
        couverts = set()
        for i, creneau in enumerate(cadre["creneaux"], start=1):
            ligne = [f"{creneau['debut']}-{creneau['fin']}"] + [""] * len(jours)
            if creneau["pause"]:
                ligne[1] = p("Pause", STYLE_PETIT)
                styles += [("SPAN", (1, i), (len(jours), i)), ("BACKGROUND", (0, i), (-1, i), colors.HexColor("#eef1f3"))]
            for c, jour in enumerate(jours, start=1):
                if (c, i) in couverts:
                    continue
                occupants = par_depart.get((jour["id"], creneau["index"]), [])
                if not occupants:
                    continue
                ligne[c] = [multiligne(*_lignes_seance(s, mode), style=STYLE_PETIT) for s, _ in occupants]
                longueur = max(l for _, l in occupants)
                libre = len(occupants) == 1 and longueur > 1 and not any(
                    par_depart.get((jour["id"], creneau["index"] + k)) for k in range(1, longueur))
                if libre:
                    styles.append(("SPAN", (c, i), (c, i + longueur - 1)))
                    couverts.update((c, i + k) for k in range(1, longueur))
                styles.append(("BACKGROUND", (c, i), (c, i + (longueur - 1 if libre else 0)), colors.HexColor("#e8f1f4")))
                if any(s["conflits"] for s, _ in occupants):
                    styles.append(("BACKGROUND", (c, i), (c, i), colors.HexColor("#f8d7d7")))
            rows.append(ligne)
        largeur = (27.5 - 2.4) / max(len(jours), 1)
        tableau = Table(rows, colWidths=[2.4 * cm] + [largeur * cm] * len(jours), repeatRows=1,
                        rowHeights=[0.7 * cm] + [1.55 * cm] * len(cadre["creneaux"]))
        style = style_tableau()
        for commande in styles:
            style.add(*commande)
        tableau.setStyle(style)
        elements += entete_etablissement() + [p(titre, STYLE_TITRE), tableau, Spacer(1, 6)]
        if donnees["emploi"]["statut"] != "PUBLIE":
            elements.append(p(f"Emploi du temps {donnees['emploi']['statut'].lower()} : non publie.", STYLE_PETIT))
    if not elements:
        elements = [p("Aucune seance a imprimer.")]
    return construire_pdf(elements, paysage=True, titre="Emploi du temps")


def xlsx_emploi(emploi, filtres, toutes=None):
    classeur = openpyxl.Workbook()
    classeur.remove(classeur.active)
    for titre, mode, donnees in _pages(emploi, filtres, toutes) or []:
        feuille = classeur.create_sheet(title="".join(c for c in titre.split(" - ", 1)[-1] if c not in "[]:*?/\\")[:31] or "Emploi")
        cadre = donnees["cadre"]
        jours = cadre["jours"]
        feuille.append([titre])
        feuille["A1"].font = Font(bold=True, size=12)
        feuille.append([])
        feuille.append(["Horaire"] + [j["nom"] for j in jours])
        for colonne in range(1, len(jours) + 2):
            cellule = feuille.cell(row=3, column=colonne)
            cellule.font = Font(bold=True)
            cellule.fill = PatternFill("solid", fgColor="EEF1F3")
        par_depart = _index_creneaux(donnees)
        for i, creneau in enumerate(cadre["creneaux"]):
            rang = 4 + i
            feuille.cell(row=rang, column=1, value=f"{creneau['debut']}-{creneau['fin']}")
            if creneau["pause"]:
                feuille.cell(row=rang, column=2, value="Pause")
                feuille.merge_cells(start_row=rang, start_column=2, end_row=rang, end_column=len(jours) + 1)
                feuille.cell(row=rang, column=2).fill = PatternFill("solid", fgColor="EEF1F3")
                continue
            for c, jour in enumerate(jours, start=2):
                occupants = par_depart.get((jour["id"], creneau["index"]), [])
                if occupants:
                    cellule = feuille.cell(row=rang, column=c, value="\n---\n".join("\n".join(_lignes_seance(s, mode)) for s, _ in occupants))
                    cellule.alignment = Alignment(wrap_text=True, vertical="top")
                    cellule.fill = PatternFill("solid", fgColor="E8F1F4")
            feuille.row_dimensions[rang].height = 48
        feuille.column_dimensions["A"].width = 14
        for c in range(2, len(jours) + 2):
            feuille.column_dimensions[get_column_letter(c)].width = 26
    if not classeur.sheetnames:
        classeur.create_sheet("Emploi").append(["Aucune seance"])
    return classeur
