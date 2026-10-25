"""Export generique des listes (CSV, Excel, PDF) : memes colonnes que l'ecran,
valeurs formatees pour la lecture (dates jj/mm/aaaa, Oui/Non)."""

import csv
import io
import re

import openpyxl
from django.http import HttpResponse
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib.units import cm
from reportlab.platypus import Spacer

from .excel import TYPE_XLSX, neutraliser_formule, orienter_feuille
from .traduction import traduire
from .pdf import Tableau as Table, STYLE_PETIT, STYLE_TITRE, construire_pdf, entete_etablissement, p, style_tableau

LIMITE_LIGNES = 20000
LIMITE_LIGNES_PDF = 2000
RE_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
RE_DATE_HEURE = re.compile(r"^\d{4}-\d{2}-\d{2}T\d{2}:\d{2}")


def formater(valeur):
    """Valeur d'API -> texte lisible."""
    if valeur is None:
        return ""
    if isinstance(valeur, bool):
        return traduire("Oui" if valeur else "Non")
    if isinstance(valeur, (list, tuple)):
        return ", ".join(formater(v) for v in valeur)
    if isinstance(valeur, dict):
        return ", ".join(f"{k}: {formater(v)}" for k, v in valeur.items())
    texte = str(valeur)
    if RE_DATE.match(texte):
        return f"{texte[8:10]}/{texte[5:7]}/{texte[:4]}"
    if RE_DATE_HEURE.match(texte):
        return f"{texte[8:10]}/{texte[5:7]}/{texte[:4]} {texte[11:16]}"
    return texte


# Codes renvoyes par l'API -> libelle (traduit ensuite selon la langue) : un export doit etre lisible sans connaitre les codes.
LIBELLES_CODES = {
    "BROUILLON": "Brouillon", "VALIDE": "Valide", "ANNULE": "Annule", "ANNULEE": "Annulee", "EN_ATTENTE": "En attente",
    "EN_COURS": "En cours", "TERMINEE": "Terminee", "LIVREE": "Livree", "OUVERTE": "Ouverte", "FERMEE": "Fermee",
    "ACTIVE": "Active", "TRANSFEREE": "Transferee", "ABANDON": "Abandon", "DIPLOME": "Diplome", "ACTIF": "Actif",
    "INACTIF": "Inactif", "PRESENT": "Present", "ABSENT": "Absent", "RETARD": "Retard", "PLANIFIE": "Planifie",
    "EFFECTUE": "Effectue", "REMPLACE": "Remplace", "ENVOYE": "Envoye", "LIVRE": "Livre", "ECHEC": "Echec",
    "NON_PAYE": "Non paye", "PARTIEL": "Partiel", "PAYE": "Paye", "PUBLIE": "Publie", "ARCHIVE": "Archive",
    "ADMINISTRATEUR": "Administrateur", "DIRECTEUR": "Directeur", "COMPTABLE": "Comptable", "PROFESSEUR": "Professeur",
    "RESPONSABLE_SCOLARITE": "Responsable de la scolarite", "RESPONSABLE_PEDAGOGIQUE": "Responsable pedagogique",
    "HORAIRE": "Au nombre d'heures", "FIXE": "Salaire fixe",
}
LIBELLES_SEXE = {"M": "Masculin", "F": "Feminin"}


def _libelle(cle, valeur):
    if valeur in LIBELLES_CODES:
        return traduire(LIBELLES_CODES[valeur])
    if cle == "sexe" and valeur in LIBELLES_SEXE:
        return traduire(LIBELLES_SEXE[valeur])
    return valeur


def lignes_de(donnees, colonnes):
    """donnees : liste de dicts (serializer) ; colonnes : [(cle, titre)] -> liste de listes de texte."""
    return [[_libelle(cle, formater(ligne.get(cle))) for cle, _ in colonnes] for ligne in donnees]


def reponse_csv(titres, lignes, nom):
    tampon = io.StringIO()
    ecrivain = csv.writer(tampon, delimiter=";")   # separateur attendu par Excel en version francaise
    ecrivain.writerow([neutraliser_formule(t) for t in titres])
    for ligne in lignes:
        ecrivain.writerow([neutraliser_formule(v) for v in ligne])
    reponse = HttpResponse(tampon.getvalue().encode("utf-8-sig"), content_type="text/csv; charset=utf-8")
    reponse["Content-Disposition"] = f'attachment; filename="{nom}.csv"'
    return reponse


def reponse_xlsx_liste(titre, titres, lignes, nom):
    classeur = openpyxl.Workbook()
    feuille = classeur.active
    feuille.title = (titre or traduire("Export"))[:31]
    orienter_feuille(feuille)
    feuille.append(titres)
    for index in range(1, len(titres) + 1):
        cellule = feuille.cell(row=1, column=index)
        cellule.font = Font(bold=True)
        cellule.fill = PatternFill("solid", fgColor="EEF1F3")
        feuille.column_dimensions[get_column_letter(index)].width = max(14, min(40, len(str(titres[index - 1])) + 6))
    for ligne in lignes:
        feuille.append([neutraliser_formule(v) for v in ligne])
    feuille.freeze_panes = "A2"
    tampon = io.BytesIO()
    classeur.save(tampon)
    reponse = HttpResponse(tampon.getvalue(), content_type=TYPE_XLSX)
    reponse["Content-Disposition"] = f'attachment; filename="{nom}.xlsx"'
    return reponse


def reponse_pdf_liste(titre, titres, lignes, nom):
    largeur_utile = 27.5 * cm
    nb = max(len(titres), 1)
    donnees = [titres] + [[p(v) for v in ligne] for ligne in lignes[:LIMITE_LIGNES_PDF]]
    tableau = Table(donnees, repeatRows=1, colWidths=[largeur_utile / nb] * nb)
    tableau.setStyle(style_tableau())
    elements = entete_etablissement(largeur_utile) + [p(titre or "Liste", STYLE_TITRE), tableau, Spacer(1, 6)]
    pied = traduire(f"{len(lignes)} ligne(s)")
    if len(lignes) > LIMITE_LIGNES_PDF:
        pied += " " + traduire(f"- PDF limite aux {LIMITE_LIGNES_PDF} premieres : utilisez l'export Excel ou CSV pour la totalite")
    elements.append(p(pied, STYLE_PETIT))
    reponse = HttpResponse(construire_pdf(elements, paysage=True, titre=titre or "Liste"), content_type="application/pdf")
    reponse["Content-Disposition"] = f'attachment; filename="{nom}.pdf"'
    return reponse
