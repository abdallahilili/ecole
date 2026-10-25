import io
from datetime import date, datetime

import openpyxl
from openpyxl.utils import get_column_letter
from django.db import transaction
from django.http import HttpResponse
from rest_framework.decorators import action
from rest_framework.response import Response
from rest_framework.parsers import MultiPartParser

from apps.commun.constantes import Role
from apps.commun.excel import orienter_feuille
from apps.commun.traduction import traduire
from apps.commun.vues_base import ViewSetGestion
from apps.referentiels.models import Classe
from .models import Etudiant, Inscription
from .serializers import EtudiantSerializer, InscriptionSerializer

COLONNES_IMPORT_ETUDIANT = [
    ("matricule", "Matricule", True),
    ("nom", "Nom", True),
    ("prenom", "Prenom", True),
    ("sexe", "Sexe (M/F)", False),
    ("date_naissance", "Date de naissance (AAAA-MM-JJ)", False),
    ("lieu_naissance", "Lieu de naissance", False),
    ("numero_national", "Numero national", False),
    ("adresse", "Adresse", False),
    ("nom_pere", "Nom du pere", False),
    ("telephone_pere", "Telephone du pere", False),
    ("nom_mere", "Nom de la mere", False),
    ("telephone_mere", "Telephone de la mere", False),
    ("nom_tuteur", "Nom du tuteur", False),
    ("telephone_tuteur", "Telephone du tuteur", False),
    ("telephone_principal", "Telephone principal (notifications)", False),
]

ROLES_LECTURE = [
    Role.ADMINISTRATEUR, Role.DIRECTEUR, Role.RESPONSABLE_SCOLARITE,
    Role.RESPONSABLE_PEDAGOGIQUE, Role.COMPTABLE, Role.PROFESSEUR,
]
ROLES_ECRITURE = [Role.ADMINISTRATEUR, Role.DIRECTEUR, Role.RESPONSABLE_SCOLARITE]


class EtudiantViewSet(ViewSetGestion):
    queryset = Etudiant.objects.all()
    serializer_class = EtudiantSerializer
    roles_lecture = ROLES_LECTURE
    roles_ecriture = ROLES_ECRITURE
    search_fields = ["matricule", "nom", "prenom", "numero_national"]
    alias_tri = {"nom_complet": ("nom", "prenom")}

    def get_queryset(self):
        qs = super().get_queryset()
        actif = self.request.query_params.get("actif")
        if actif is not None:
            qs = qs.filter(actif=actif.lower() in ("1", "true", "vrai"))
        # ?classe=<id> : eleves inscrits (inscription active) dans cette classe ; une classe appartient a une seule annee.
        classe = self.request.query_params.get("classe")
        if classe:
            qs = qs.filter(inscriptions__classe_id=classe, inscriptions__statut="ACTIVE").distinct()
        return qs

    @action(detail=False, methods=["get"], url_path="modele-import")
    def modele_import(self, request):
        """Modele Excel vierge a telecharger et completer pour l'import."""
        classeur = openpyxl.Workbook()
        feuille = classeur.active
        feuille.title = traduire("Etudiants")
        orienter_feuille(feuille)
        for index, (_, titre, _obligatoire) in enumerate(COLONNES_IMPORT_ETUDIANT, start=1):
            titre = traduire(titre)
            feuille.cell(row=1, column=index, value=titre)
            feuille.column_dimensions[get_column_letter(index)].width = max(18, len(titre) + 2)
        feuille.cell(row=2, column=1, value="E1001")
        feuille.cell(row=2, column=2, value="Diallo")
        feuille.cell(row=2, column=3, value="Awa")
        feuille.cell(row=2, column=4, value="F")

        tampon = io.BytesIO()
        classeur.save(tampon)
        tampon.seek(0)
        reponse = HttpResponse(tampon.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        reponse["Content-Disposition"] = 'attachment; filename="modele_import_etudiants.xlsx"'
        return reponse

    @action(detail=False, methods=["get"], url_path="exporter")
    def exporter(self, request):
        """Exporte la liste (filtree) des etudiants en Excel."""
        qs = self.filter_queryset(self.get_queryset())
        classeur = openpyxl.Workbook()
        feuille = classeur.active
        feuille.title = traduire("Etudiants")
        orienter_feuille(feuille)
        entetes = [traduire(h) for h in ["Matricule", "Nom", "Prenom", "Sexe", "Date de naissance", "Telephone principal", "Actif"]]
        feuille.append(entetes)
        for etudiant in qs:
            feuille.append([
                etudiant.matricule, etudiant.nom, etudiant.prenom, etudiant.sexe,
                etudiant.date_naissance.isoformat() if etudiant.date_naissance else "",
                etudiant.telephone_notification, traduire("Oui" if etudiant.actif else "Non"),
            ])
        for index in range(1, len(entetes) + 1):
            feuille.column_dimensions[get_column_letter(index)].width = 20

        tampon = io.BytesIO()
        classeur.save(tampon)
        tampon.seek(0)
        reponse = HttpResponse(tampon.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        reponse["Content-Disposition"] = 'attachment; filename="etudiants.xlsx"'
        return reponse

    @action(detail=False, methods=["post"], url_path="importer", parser_classes=[MultiPartParser])
    def importer(self, request):
        """Import Excel valide ligne par ligne AVANT tout enregistrement.
        Si une seule ligne est invalide, rien n'est enregistre (jamais
        d'ecrasement silencieux ni d'import partiel)."""
        fichier = request.FILES.get("fichier")
        if not fichier:
            return Response({"detail": "Aucun fichier recu (champ 'fichier' attendu)."}, status=400)

        try:
            classeur = openpyxl.load_workbook(fichier, data_only=True)
        except Exception:
            return Response({"detail": "Fichier Excel illisible ou corrompu."}, status=400)
        feuille = classeur.active

        cles = [cle for cle, _titre, _obligatoire in COLONNES_IMPORT_ETUDIANT]
        obligatoires = {cle for cle, _titre, obligatoire in COLONNES_IMPORT_ETUDIANT if obligatoire}

        lignes = list(feuille.iter_rows(min_row=2, values_only=True))
        erreurs = []
        a_creer = []
        matricules_fichier = set()
        matricules_existants = set(Etudiant.objects.values_list("matricule", flat=True))

        for numero, ligne in enumerate(lignes, start=2):
            if ligne is None or all(valeur in (None, "") for valeur in ligne):
                continue
            donnees = {}
            for index, cle in enumerate(cles):
                valeur = ligne[index] if index < len(ligne) else None
                donnees[cle] = valeur

            erreurs_ligne = []
            for cle in obligatoires:
                if not donnees.get(cle) or not str(donnees.get(cle)).strip():
                    erreurs_ligne.append(f"colonne '{cle}' obligatoire manquante")

            matricule = str(donnees.get("matricule") or "").strip()
            if matricule:
                if matricule in matricules_existants:
                    erreurs_ligne.append(f"matricule '{matricule}' existe deja")
                if matricule in matricules_fichier:
                    erreurs_ligne.append(f"matricule '{matricule}' duplique dans le fichier")
                matricules_fichier.add(matricule)

            date_naissance = donnees.get("date_naissance")
            if isinstance(date_naissance, str) and date_naissance.strip():
                try:
                    date_naissance = datetime.strptime(date_naissance.strip(), "%Y-%m-%d").date()
                except ValueError:
                    erreurs_ligne.append("date de naissance invalide (attendu AAAA-MM-JJ)")
                    date_naissance = None
            elif hasattr(date_naissance, "date"):
                date_naissance = date_naissance.date()

            sexe = str(donnees.get("sexe") or "M").strip().upper()[:1] or "M"
            if sexe not in ("M", "F"):
                erreurs_ligne.append("sexe invalide (attendu M ou F)")

            if erreurs_ligne:
                erreurs.append({"ligne": numero, "erreurs": erreurs_ligne})
                continue

            a_creer.append(Etudiant(
                matricule=matricule, nom=str(donnees.get("nom") or "").strip(),
                prenom=str(donnees.get("prenom") or "").strip(), sexe=sexe,
                date_naissance=date_naissance,
                lieu_naissance=str(donnees.get("lieu_naissance") or "").strip(),
                numero_national=str(donnees.get("numero_national") or "").strip(),
                adresse=str(donnees.get("adresse") or "").strip(),
                nom_pere=str(donnees.get("nom_pere") or "").strip(),
                telephone_pere=str(donnees.get("telephone_pere") or "").strip(),
                nom_mere=str(donnees.get("nom_mere") or "").strip(),
                telephone_mere=str(donnees.get("telephone_mere") or "").strip(),
                nom_tuteur=str(donnees.get("nom_tuteur") or "").strip(),
                telephone_tuteur=str(donnees.get("telephone_tuteur") or "").strip(),
                telephone_principal=str(donnees.get("telephone_principal") or "").strip(),
            ))

        if erreurs:
            return Response({"detail": "Import refuse : des lignes contiennent des erreurs.", "erreurs": erreurs}, status=400)
        if not a_creer:
            return Response({"detail": "Le fichier ne contient aucune ligne a importer."}, status=400)

        # ?classe=<id> : les eleves importes sont inscrits dans cette classe (import depuis la liste d'une classe).
        classe = None
        if request.query_params.get("classe"):
            classe = Classe.objects.filter(pk=request.query_params["classe"]).select_related("annee_scolaire").first()
            if classe is None:
                return Response({"detail": "Classe introuvable."}, status=400)
            effectif = Inscription.objects.filter(classe=classe, statut="ACTIVE").count()
            if effectif + len(a_creer) > classe.capacite:
                return Response({"detail": f"La classe '{classe.nom}' a atteint sa capacite maximale ({classe.capacite} eleves)."}, status=400)

        with transaction.atomic():
            Etudiant.objects.bulk_create(a_creer)
            if classe is not None:
                crees = Etudiant.objects.filter(matricule__in=[e.matricule for e in a_creer])
                Inscription.objects.bulk_create([
                    Inscription(etudiant=e, annee_scolaire=classe.annee_scolaire, classe=classe, date_inscription=date.today())
                    for e in crees
                ])
        return Response({"detail": f"{len(a_creer)} etudiant(s) importe(s) avec succes.", "nombre_importes": len(a_creer)}, status=201)


class InscriptionViewSet(ViewSetGestion):
    filtre_actif = ("statut", "ACTIVE")
    queryset = Inscription.objects.select_related("etudiant", "classe", "groupe", "annee_scolaire").all()
    serializer_class = InscriptionSerializer
    roles_lecture = ROLES_LECTURE
    roles_ecriture = ROLES_ECRITURE

    def get_queryset(self):
        qs = super().get_queryset()
        etudiant = self.request.query_params.get("etudiant")
        if etudiant:
            qs = qs.filter(etudiant_id=etudiant)
        classe = self.request.query_params.get("classe")
        if classe:
            qs = qs.filter(classe_id=classe)
        groupe = self.request.query_params.get("groupe")
        if groupe:
            qs = qs.filter(groupe_id=groupe)
        annee = self.request.query_params.get("annee_scolaire")
        if annee:
            qs = qs.filter(annee_scolaire_id=annee)
        statut = self.request.query_params.get("statut")
        if statut:
            qs = qs.filter(statut=statut)
        return qs
