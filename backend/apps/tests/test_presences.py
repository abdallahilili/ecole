import hashlib
import hmac
import io
import json
from datetime import date, time
from decimal import Decimal
from unittest.mock import patch

import openpyxl
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from rest_framework.test import APIClient

from apps.commun.constantes import Role
from apps.commun.models import JournalAudit, Parametre
from apps.etudiants.models import Etudiant, Inscription
from apps.presences import whatsapp
from apps.presences.models import (
    NotificationWhatsApp, PresenceEtudiant, Seance, StatutNotification, StatutPresence,
)
from apps.presences.services import normaliser_telephone
from apps.professeurs.models import AffectationProfesseur, Professeur
from apps.referentiels.models import AnneeScolaire, Classe, Matiere, Niveau
from apps.utilisateurs.models import Utilisateur

API_WHATSAPP = dict(WHATSAPP_JETON="jeton", WHATSAPP_ID_NUMERO="123", WHATSAPP_MODELE="absence_eleve")
SEANCES = "/api/presences/seances/"


def utilisateur(role=Role.ADMINISTRATEUR, identifiant="admin", **extra):
    return Utilisateur.objects.create_user(identifiant=identifiant, mot_de_passe="secret123", nom="Test", role=role, **extra)


def client_de(user):
    client = APIClient()
    client.force_authenticate(user)
    return client


class BasePresences(TestCase):
    def setUp(self):
        self.annee = AnneeScolaire.objects.create(
            libelle="2025-2026", date_debut=date(2025, 9, 1), date_fin=date(2026, 6, 30), active=True
        )
        self.niveau = Niveau.objects.create(nom="6eme", code="6E", ordre=1)
        self.classe = Classe.objects.create(nom="6AF", niveau=self.niveau, annee_scolaire=self.annee, capacite=30)
        self.autre_classe = Classe.objects.create(nom="7D", niveau=self.niveau, annee_scolaire=self.annee, capacite=30)
        self.matiere = Matiere.objects.create(nom="Mathematiques", code="MATH")
        self.prof = Professeur.objects.create(matricule="P1", nom="Sy", prenom="Oumar", tarif_horaire=100)
        self.prof2 = Professeur.objects.create(matricule="P2", nom="Ba", prenom="Cheikh", tarif_horaire=100)
        self.awa = self.eleve("E1", "Diallo", "Awa", classe=self.classe, telephone_principal="22288896")
        self.moussa = self.eleve("E2", "Ba", "Moussa", classe=self.classe, telephone_pere="31177008", nom_pere="Ba Amadou")
        self.sans_tel = self.eleve("E3", "Fall", "Mariam", classe=self.classe)
        self.autre = self.eleve("E9", "Sow", "Ali", classe=self.autre_classe)
        self.admin = utilisateur()
        self.client = client_de(self.admin)

    def eleve(self, matricule, nom, prenom, classe, **extra):
        etudiant = Etudiant.objects.create(matricule=matricule, nom=nom, prenom=prenom, **extra)
        Inscription.objects.create(etudiant=etudiant, annee_scolaire=self.annee, classe=classe, date_inscription=date(2025, 9, 1))
        return etudiant

    def cours(self, **extra):
        donnees = {
            "annee_scolaire": self.annee.id, "classe": self.classe.id, "matiere": self.matiere.id,
            "date": "2025-10-06", "heure_debut": "08:00", "heure_fin": "09:00",
        }
        donnees.update(extra)
        return donnees

    def seance(self, **extra):
        defauts = dict(
            annee_scolaire=self.annee, classe=self.classe, matiere=self.matiere, professeur=self.prof,
            date=date(2025, 10, 6), heure_debut=time(8, 0), heure_fin=time(9, 0),
        )
        defauts.update(extra)
        return Seance.objects.create(**defauts)

    def marquer(self, seance, statuts):
        """statuts : {etudiant: statut}"""
        for etudiant, statut in statuts.items():
            PresenceEtudiant.objects.create(seance=seance, etudiant=etudiant, statut=statut)

    def tout_marquer(self, seance, **statuts):
        defauts = {self.awa: "PRESENT", self.moussa: "PRESENT", self.sans_tel: "PRESENT"}
        defauts.update(statuts)
        self.marquer(seance, defauts)


class HeuresProfesseurTests(BasePresences):
    def test_heures_selon_le_statut(self):
        s = self.seance(statut_professeur="EFFECTUE")
        self.assertEqual(s.heures_effectuees, Decimal("1.00"))
        s.statut_professeur = "ABSENT"
        s.save()
        self.assertEqual(s.heures_effectuees, Decimal("0.00"))
        s.statut_professeur = "ANNULE"
        s.save()
        self.assertEqual(s.heures_effectuees, Decimal("0.00"))

    def test_retard_utilise_les_heures_reelles(self):
        s = self.seance(heure_fin=time(10, 0), statut_professeur="RETARD", heure_reelle_debut=time(8, 30))
        self.assertEqual(s.heures_effectuees, Decimal("1.50"))

    def test_remplacement_credite_le_remplacant(self):
        s = self.seance(statut_professeur="REMPLACE", professeur_remplacant=self.prof2)
        self.assertEqual(s.professeur_effectif, self.prof2)
        self.assertEqual(s.heures_effectuees, Decimal("1.00"))

    def test_api_calcule_les_heures_cote_serveur(self):
        reponse = self.client.post(SEANCES, {**self.cours(), "professeur": self.prof.id, "statut_professeur": "EFFECTUE",
                                             "heures_effectuees": "99"})
        self.assertEqual(reponse.status_code, 201, reponse.content)
        self.assertEqual(reponse.json()["heures_effectuees"], "1.00")

    def test_resume_des_heures_par_professeur_effectif(self):
        autre = Matiere.objects.create(nom="Physique", code="PHY")
        self.seance(statut_professeur="EFFECTUE")                                              # Sy : 1 h
        self.seance(matiere=autre, heure_debut=time(9, 0), heure_fin=time(11, 0), statut_professeur="REMPLACE",
                    professeur_remplacant=self.prof2)                                          # Ba : 2 h
        self.seance(heure_debut=time(11, 0), heure_fin=time(12, 0), statut_professeur="ABSENT")  # 0 h
        self.seance(heure_debut=time(13, 0), heure_fin=time(14, 0), statut_professeur="ANNULE")  # 0 h
        resume = self.client.get(SEANCES + "resume/").json()
        self.assertEqual(resume["heures_effectuees"], "3.00")
        self.assertEqual({p["nom"]: p["heures"] for p in resume["par_professeur"]}, {"Sy Oumar": "1.00", "Ba Cheikh": "2.00"})
        self.assertEqual(resume["par_statut"], {"EFFECTUE": 1, "REMPLACE": 1, "ABSENT": 1, "ANNULE": 1})
        filtre = self.client.get(SEANCES + "resume/", {"professeur": self.prof2.id}).json()
        self.assertEqual(filtre["heures_effectuees"], "2.00")

    def test_validations_du_statut(self):
        base = {**self.cours(), "professeur": self.prof.id}
        self.assertEqual(self.client.post(SEANCES, {**base, "statut_professeur": "REMPLACE"}).status_code, 400)
        self.assertEqual(self.client.post(SEANCES, {**base, "statut_professeur": "RETARD"}).status_code, 400)
        self.assertEqual(self.client.post(SEANCES, {
            **base, "statut_professeur": "REMPLACE", "professeur_remplacant": self.prof.id}).status_code, 400)


class ConflitsSeanceTests(BasePresences):
    def test_classe_ou_professeur_ne_peuvent_pas_avoir_deux_cours_simultanes(self):
        self.seance()
        autre_matiere = Matiere.objects.create(nom="Physique", code="PHY")
        # Meme classe, creneau qui chevauche.
        r = self.client.post(SEANCES, {**self.cours(matiere=autre_matiere.id, heure_debut="08:30", heure_fin="09:30")})
        self.assertEqual(r.status_code, 400)
        # Autre classe mais meme professeur.
        r = self.client.post(SEANCES, {**self.cours(classe=self.autre_classe.id, matiere=autre_matiere.id,
                                                    heure_debut="08:30", heure_fin="09:30"), "professeur": self.prof.id})
        self.assertEqual(r.status_code, 400)
        # Creneaux consecutifs : autorises.
        r = self.client.post(SEANCES, {**self.cours(matiere=autre_matiere.id, heure_debut="09:00", heure_fin="10:00")})
        self.assertEqual(r.status_code, 201, r.content)

    def test_seance_annulee_ne_bloque_pas_le_creneau(self):
        self.seance(statut_professeur="ANNULE")
        autre_matiere = Matiere.objects.create(nom="Physique", code="PHY")
        r = self.client.post(SEANCES, {**self.cours(matiere=autre_matiere.id), "professeur": self.prof.id})
        self.assertEqual(r.status_code, 201, r.content)
        # Le meme cours (classe, matiere, date, heure) ne peut exister qu'une fois.
        self.assertEqual(self.client.post(SEANCES, {**self.cours(), "professeur": self.prof2.id}).status_code, 400)

    def test_date_hors_annee_scolaire_refusee(self):
        self.assertEqual(self.client.post(SEANCES, self.cours(date="2026-09-15")).status_code, 400)

    def test_classe_d_une_autre_annee_refusee(self):
        autre = AnneeScolaire.objects.create(libelle="2026-2027", date_debut=date(2026, 9, 1), date_fin=date(2027, 6, 30))
        self.assertEqual(self.client.post(SEANCES, self.cours(annee_scolaire=autre.id, date="2026-10-05")).status_code, 400)


class ListeAppelTests(BasePresences):
    def test_ouvrir_liste_uniquement_les_eleves_de_la_classe(self):
        reponse = self.client.post(SEANCES + "ouvrir/", self.cours())
        self.assertEqual(reponse.status_code, 200, reponse.content)
        matricules = [e["matricule"] for e in reponse.json()["eleves"]]
        self.assertEqual(sorted(matricules), ["E1", "E2", "E3"])
        self.assertNotIn("E9", matricules)
        # Idempotent : pas de seance en double.
        self.client.post(SEANCES + "ouvrir/", self.cours())
        self.assertEqual(Seance.objects.count(), 1)

    def test_enregistrer_refuse_un_eleve_hors_classe(self):
        seance = self.seance()
        r = self.client.post(f"{SEANCES}{seance.id}/enregistrer/", {"presences": [{"etudiant": self.autre.id, "statut": "ABSENT"}]}, format="json")
        self.assertEqual(r.status_code, 400)
        self.assertEqual(PresenceEtudiant.objects.count(), 0)

    def test_enregistrer_puis_effacer(self):
        seance = self.seance()
        url = f"{SEANCES}{seance.id}/enregistrer/"
        self.client.post(url, {"presences": [{"etudiant": self.awa.id, "statut": "ABSENT", "observation": "malade"}]}, format="json")
        self.assertEqual(PresenceEtudiant.objects.get(etudiant=self.awa).observation, "malade")
        self.client.post(url, {"presences": [{"etudiant": self.awa.id, "statut": ""}]}, format="json")
        self.assertEqual(PresenceEtudiant.objects.count(), 0)
        r = self.client.post(url, {"presences": [{"etudiant": self.awa.id, "statut": "ABSENCE"}]}, format="json")
        self.assertEqual(r.status_code, 400)

    def test_suppression_refusee_si_presences(self):
        seance = self.seance()
        self.assertEqual(self.client.delete(f"{SEANCES}{seance.id}/").status_code, 204)
        seance = self.seance()
        self.marquer(seance, {self.awa: "PRESENT"})
        self.assertEqual(self.client.delete(f"{SEANCES}{seance.id}/").status_code, 400)


class ValidationEtNotificationsTests(BasePresences):
    def activer_whatsapp(self, retard_absence=False):
        Parametre.objects.update_or_create(cle="whatsapp_actif", defaults={"valeur": "1"})
        Parametre.objects.update_or_create(cle="retard_est_absence", defaults={"valeur": "1" if retard_absence else "0"})

    def valider(self, seance):
        return self.client.post(f"{SEANCES}{seance.id}/valider/")

    def test_validation_exige_tous_les_statuts(self):
        seance = self.seance()
        self.marquer(seance, {self.awa: "PRESENT"})
        r = self.valider(seance)
        self.assertEqual(r.status_code, 400)
        self.assertIn("sans statut", r.json()["detail"])
        seance.refresh_from_db()
        self.assertFalse(seance.presences_validees)

    def test_notifie_uniquement_les_absents(self):
        self.activer_whatsapp()
        seance = self.seance()
        self.marquer(seance, {self.awa: "ABSENT", self.moussa: "PRESENT", self.sans_tel: "RETARD"})
        r = self.valider(seance)
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual(r.json()["notifications"]["creees"], 1)
        notification = NotificationWhatsApp.objects.get()
        self.assertEqual(notification.presence.etudiant, self.awa)
        self.assertEqual(notification.telephone, "22222288896")
        self.assertEqual(notification.statut, StatutNotification.EN_ATTENTE)
        for attendu in ("Diallo Awa", "6AF", "Mathematiques", "06/10/2025", "08:00", "09:00"):
            self.assertIn(attendu, notification.message)

    def test_retard_notifie_seulement_si_assimile_a_une_absence(self):
        self.activer_whatsapp(retard_absence=True)
        seance = self.seance()
        self.marquer(seance, {self.awa: "RETARD", self.moussa: "PRESENT", self.sans_tel: "PRESENT"})
        self.valider(seance)
        self.assertEqual(NotificationWhatsApp.objects.count(), 1)

    def test_envoi_desactive_aucune_notification(self):
        seance = self.seance()
        self.marquer(seance, {self.awa: "ABSENT", self.moussa: "PRESENT", self.sans_tel: "PRESENT"})
        r = self.valider(seance)
        self.assertEqual(r.status_code, 200)
        self.assertFalse(r.json()["notifications"]["actif"])
        self.assertEqual(NotificationWhatsApp.objects.count(), 0)

    def test_absent_sans_telephone_est_en_echec_visible(self):
        self.activer_whatsapp()
        seance = self.seance()
        self.marquer(seance, {self.awa: "PRESENT", self.moussa: "PRESENT", self.sans_tel: "ABSENT"})
        r = self.valider(seance)
        self.assertEqual(r.json()["notifications"]["sans_telephone"], ["Fall Mariam"])
        self.assertEqual(NotificationWhatsApp.objects.get().statut, StatutNotification.ECHEC)

    def test_seance_validee_verrouillee_et_pas_de_doublon_apres_reouverture(self):
        self.activer_whatsapp()
        seance = self.seance()
        self.marquer(seance, {self.awa: "ABSENT", self.moussa: "PRESENT", self.sans_tel: "PRESENT"})
        self.valider(seance)
        url = f"{SEANCES}{seance.id}/"
        self.assertEqual(self.client.post(url + "enregistrer/", {"presences": []}, format="json").status_code, 400)
        self.assertEqual(self.client.post(url + "valider/").status_code, 400)
        self.assertEqual(self.client.patch(url, {"heure_fin": "10:00"}).status_code, 400)
        # Reouverture : motif obligatoire, puis revalidation sans doublon.
        self.assertEqual(self.client.post(url + "rouvrir/", {}).status_code, 400)
        self.assertEqual(self.client.post(url + "rouvrir/", {"motif": "erreur de saisie"}).status_code, 200)
        self.assertTrue(JournalAudit.objects.filter(action="REOUVERTURE_PRESENCES").exists())
        self.assertEqual(self.valider(seance).json()["notifications"]["creees"], 0)
        self.assertEqual(NotificationWhatsApp.objects.count(), 1)

    def test_validation_tracee_dans_l_audit(self):
        seance = self.seance()
        self.tout_marquer(seance)
        self.valider(seance)
        self.assertTrue(JournalAudit.objects.filter(action="VALIDATION_PRESENCES", reference_objet=str(seance.pk)).exists())

    def test_notification_manuelle_sans_doublon(self):
        self.activer_whatsapp()
        seance = self.seance()
        self.marquer(seance, {self.awa: "PRESENT", self.moussa: "PRESENT", self.sans_tel: "PRESENT"})
        presence = PresenceEtudiant.objects.get(etudiant=self.awa)
        url = "/api/presences/notifications/manuelle/"
        self.assertEqual(self.client.post(url, {"presence": presence.id}).status_code, 400)  # non validee
        self.valider(seance)
        self.assertEqual(self.client.post(url, {"presence": presence.id}).status_code, 400)  # present
        presence.statut = StatutPresence.ABSENT
        presence.save()
        self.assertEqual(self.client.post(url, {"presence": presence.id}).status_code, 201)
        self.assertEqual(self.client.post(url, {"presence": presence.id}).status_code, 400)  # doublon


class EnvoiWhatsAppTests(BasePresences):
    def notification(self, **extra):
        seance = self.seance()
        presence = PresenceEtudiant.objects.create(seance=seance, etudiant=self.awa, statut="ABSENT")
        return NotificationWhatsApp.objects.create(presence=presence, telephone="22222288896", message="Bonjour", **extra)

    def test_normaliser_telephone(self):
        self.assertEqual(normaliser_telephone("22288896"), "22222288896")
        self.assertEqual(normaliser_telephone("+222 22 28 88 96"), "22222288896")
        self.assertEqual(normaliser_telephone("00222 31177008"), "22231177008")
        self.assertIsNone(normaliser_telephone("123"))
        self.assertIsNone(normaliser_telephone(""))

    def test_sans_configuration_rien_n_est_envoye(self):
        notification = self.notification()
        with patch("apps.presences.whatsapp.envoyer_modele") as envoi:
            bilan = whatsapp.traiter_file()
        envoi.assert_not_called()
        self.assertFalse(bilan["api_configuree"])
        notification.refresh_from_db()
        self.assertEqual(notification.statut, StatutNotification.EN_ATTENTE)

    @override_settings(**API_WHATSAPP)
    def test_envoye_seulement_apres_acceptation_par_l_api(self):
        notification = self.notification()
        with patch("apps.presences.whatsapp.envoyer_modele", return_value="wamid.ABC") as envoi:
            bilan = whatsapp.traiter_file()
        self.assertEqual(bilan["envoyees"], 1)
        # 7 variables du modele approuve, dans l'ordre.
        self.assertEqual(envoi.call_args[0][1][:3], ["Diallo Awa", "6AF", "Mathematiques"])
        self.assertEqual(len(envoi.call_args[0][1]), 7)
        notification.refresh_from_db()
        self.assertEqual((notification.statut, notification.canal, notification.identifiant_message_api),
                         (StatutNotification.ENVOYE, "API", "wamid.ABC"))
        self.assertIsNotNone(notification.envoye_le)

    @override_settings(**API_WHATSAPP)
    def test_reponse_sans_identifiant_n_est_jamais_envoyee(self):
        notification = self.notification()
        reponse = io.BytesIO(json.dumps({"messages": []}).encode())
        with patch("urllib.request.urlopen") as ouvrir:
            ouvrir.return_value.__enter__.return_value = reponse
            whatsapp.traiter_file()
        notification.refresh_from_db()
        self.assertEqual(notification.statut, StatutNotification.EN_ATTENTE)
        self.assertIn("acceptation", notification.erreur)

    @override_settings(**API_WHATSAPP, WHATSAPP_MAX_TENTATIVES=2)
    def test_echec_transitoire_puis_definitif_apres_max_tentatives(self):
        notification = self.notification()
        erreur = whatsapp.ErreurWhatsApp("reseau coupe", definitif=False)
        with patch("apps.presences.whatsapp.envoyer_modele", side_effect=erreur):
            bilan = whatsapp.traiter_file()
            self.assertEqual(bilan["reportees"], 1)
            notification.refresh_from_db()
            self.assertEqual((notification.statut, notification.tentatives), (StatutNotification.EN_ATTENTE, 1))
            # Pas de nouvel essai avant l'echeance...
            self.assertEqual(whatsapp.traiter_file()["reportees"], 0)
            # ... puis nouvel essai : maximum atteint -> echec.
            from django.utils import timezone
            from datetime import timedelta
            bilan = whatsapp.traiter_file(maintenant=timezone.now() + timedelta(hours=1))
        self.assertEqual(bilan["echecs"], 1)
        notification.refresh_from_db()
        self.assertEqual(notification.statut, StatutNotification.ECHEC)

    @override_settings(**API_WHATSAPP)
    def test_echec_definitif_immediat_et_reessai_manuel(self):
        notification = self.notification()
        with patch("apps.presences.whatsapp.envoyer_modele", side_effect=whatsapp.ErreurWhatsApp("numero invalide", definitif=True)):
            whatsapp.traiter_file()
        notification.refresh_from_db()
        self.assertEqual(notification.statut, StatutNotification.ECHEC)
        with patch("apps.presences.whatsapp.envoyer_modele", return_value="wamid.OK"):
            r = self.client.post(f"/api/presences/notifications/{notification.id}/reessayer/")
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual(r.json()["statut"], "ENVOYE")

    def test_marquer_envoye_manuellement_et_lien_wa_me(self):
        notification = self.notification()
        r = self.client.get(f"/api/presences/notifications/{notification.id}/")
        self.assertTrue(r.json()["lien_wa_me"].startswith("https://wa.me/22222288896?text="))
        r = self.client.post(f"/api/presences/notifications/{notification.id}/marquer-envoye/")
        self.assertEqual((r.json()["statut"], r.json()["canal"]), ("ENVOYE", "MANUEL"))
        self.assertEqual(self.client.post(f"/api/presences/notifications/{notification.id}/marquer-envoye/").status_code, 400)

    def test_reessayer_sans_numero_valide_refuse(self):
        notification = self.notification(statut=StatutNotification.ECHEC)
        r = self.client.post(f"/api/presences/notifications/{notification.id}/reessayer/")
        self.assertEqual(r.status_code, 200)  # le numero de l'eleve (22288896) est valide
        self.awa.telephone_principal = ""
        self.awa.save()
        notification.statut = StatutNotification.ECHEC
        notification.save()
        r = self.client.post(f"/api/presences/notifications/{notification.id}/reessayer/")
        self.assertEqual(r.status_code, 400)

    @override_settings(WHATSAPP_SECRET_APPLICATION="secret", WHATSAPP_JETON_VERIFICATION="verif")
    def test_webhook_signature_et_livraison(self):
        notification = self.notification(statut=StatutNotification.ENVOYE, identifiant_message_api="wamid.1")
        corps = json.dumps({"entry": [{"changes": [{"value": {"statuses": [{"id": "wamid.1", "status": "delivered"}]}}]}]}).encode()
        signature = "sha256=" + hmac.new(b"secret", corps, hashlib.sha256).hexdigest()
        anonyme = APIClient()
        mauvaise = anonyme.post("/api/presences/whatsapp/webhook/", corps, content_type="application/json",
                                HTTP_X_HUB_SIGNATURE_256="sha256=000")
        self.assertEqual(mauvaise.status_code, 403)
        bonne = anonyme.post("/api/presences/whatsapp/webhook/", corps, content_type="application/json",
                             HTTP_X_HUB_SIGNATURE_256=signature)
        self.assertEqual(bonne.status_code, 200)
        notification.refresh_from_db()
        self.assertEqual(notification.statut, StatutNotification.LIVRE)
        # Verification de l'abonnement par Meta.
        ok = anonyme.get("/api/presences/whatsapp/webhook/", {"hub.mode": "subscribe", "hub.verify_token": "verif", "hub.challenge": "42"})
        self.assertEqual((ok.status_code, ok.content), (200, b"42"))
        self.assertEqual(anonyme.get("/api/presences/whatsapp/webhook/", {"hub.mode": "subscribe", "hub.verify_token": "x"}).status_code, 403)

    def test_webhook_refuse_sans_secret_configure(self):
        r = APIClient().post("/api/presences/whatsapp/webhook/", b"{}", content_type="application/json",
                             HTTP_X_HUB_SIGNATURE_256="sha256=abc")
        self.assertEqual(r.status_code, 403)

    def test_configuration_n_expose_aucun_secret(self):
        with override_settings(**API_WHATSAPP):
            r = self.client.get("/api/presences/notifications/configuration/")
        self.assertTrue(r.json()["api_configuree"])
        self.assertNotIn("jeton", json.dumps(r.json()).lower().replace("nom_modele_api", ""))


class ExportsTests(BasePresences):
    def test_pdf_et_excel_de_la_liste(self):
        seance = self.seance()
        self.marquer(seance, {self.awa: "ABSENT"})
        pdf = self.client.get(f"{SEANCES}{seance.id}/pdf/")
        self.assertEqual(pdf["Content-Type"], "application/pdf")
        self.assertTrue(pdf.content.startswith(b"%PDF"))
        excel = self.client.get(f"{SEANCES}{seance.id}/excel/")
        feuille = openpyxl.load_workbook(io.BytesIO(excel.content)).active
        valeurs = [c.value for ligne in feuille.iter_rows() for c in ligne]
        self.assertIn("Diallo Awa", valeurs)
        self.assertNotIn("Sow Ali", valeurs)

    def test_modele_vierge(self):
        pdf = self.client.get(SEANCES + "modele-vierge/pdf/")
        self.assertTrue(pdf.content.startswith(b"%PDF"))
        feuille = openpyxl.load_workbook(io.BytesIO(self.client.get(SEANCES + "modele-vierge/excel/").content)).active
        self.assertGreater(feuille.max_row, 40)


class ImportExcelTests(BasePresences):
    def modele(self):
        r = self.client.get(SEANCES + "modele-import/", self.cours())
        self.assertEqual(r.status_code, 200, r.content)
        return openpyxl.load_workbook(io.BytesIO(r.content))

    def fichier(self, classeur):
        tampon = io.BytesIO()
        classeur.save(tampon)
        return SimpleUploadedFile("presences.xlsx", tampon.getvalue())

    def remplir(self, classeur, statuts):
        """statuts : {matricule: valeur}"""
        feuille = classeur.active
        for ligne in feuille.iter_rows(min_row=9):
            if ligne[1].value in statuts:
                ligne[3].value = statuts[ligne[1].value]
        return classeur

    def envoyer(self, classeur, **extra):
        return self.client.post(SEANCES + "importer/", {**self.cours(), "fichier": self.fichier(classeur), **extra}, format="multipart")

    def test_le_modele_contient_les_eleves_de_la_classe(self):
        feuille = self.modele().active
        self.assertEqual(feuille["B2"].value, "6AF")
        self.assertEqual(feuille["B3"].value, "MATH")
        matricules = {ligne[1].value for ligne in feuille.iter_rows(min_row=9) if ligne[1].value}
        self.assertEqual(matricules, {"E1", "E2", "E3"})

    def test_apercu_n_ecrit_rien_puis_confirmation_enregistre(self):
        classeur = self.remplir(self.modele(), {"E1": "A", "E2": "P", "E3": "r"})
        apercu = self.envoyer(classeur)
        self.assertEqual(apercu.status_code, 200, apercu.content)
        self.assertEqual(apercu.json()["resume"]["nouveaux"], 3)
        self.assertEqual(apercu.json()["blocages"], [])
        self.assertEqual((Seance.objects.count(), PresenceEtudiant.objects.count()), (0, 0))

        confirme = self.envoyer(classeur, confirmer="true")
        self.assertEqual(confirme.status_code, 201, confirme.content)
        statuts = {p.etudiant.matricule: p.statut for p in PresenceEtudiant.objects.all()}
        self.assertEqual(statuts, {"E1": "ABSENT", "E2": "PRESENT", "E3": "RETARD"})
        self.assertFalse(Seance.objects.get().presences_validees)  # brouillon : a valider ensuite
        self.assertTrue(JournalAudit.objects.filter(action="IMPORT_PRESENCES").exists())

    def test_anomalies_detectees_et_rien_enregistre(self):
        classeur = self.remplir(self.modele(), {"E1": "A", "E2": "Z", "E3": None})
        feuille = classeur.active
        feuille.append([None, "E9", "Sow Ali", "P", None])      # etudiant d'une autre classe
        feuille.append([None, "XXXX", "Inconnu", "P", None])    # matricule inconnu
        feuille.append([None, "E1", "Diallo Awa", "A", None])   # doublon
        resultat = self.envoyer(classeur, confirmer="true")
        self.assertEqual(resultat.status_code, 400)
        anomalies = {l["matricule"]: " ".join(l["anomalies"]) for l in resultat.json()["lignes"] if l["anomalies"]}
        self.assertIn("statut invalide", anomalies["E2"])
        self.assertIn("statut manquant", anomalies["E3"])
        self.assertIn("hors de cette classe", anomalies["E9"])
        self.assertIn("inconnu", anomalies["XXXX"])
        self.assertIn("double", anomalies["E1"])
        self.assertEqual(PresenceEtudiant.objects.count(), 0)

    def test_fichier_d_un_autre_cours_refuse(self):
        classeur = self.remplir(self.modele(), {"E1": "A", "E2": "P", "E3": "P"})
        classeur.active["B4"].value = "2025-10-07"
        resultat = self.envoyer(classeur, confirmer="true")
        self.assertEqual(resultat.status_code, 400)
        self.assertIn("date du fichier", resultat.json()["detail"])

    def test_aucun_ecrasement_silencieux(self):
        seance = self.seance()
        self.marquer(seance, {self.awa: "PRESENT"})
        classeur = self.remplir(self.modele(), {"E1": "A", "E2": "P", "E3": "P"})
        apercu = self.envoyer(classeur).json()
        self.assertEqual(apercu["resume"]["modifies"], 1)
        self.assertTrue(apercu["blocages"])
        refuse = self.envoyer(classeur, confirmer="true")
        self.assertEqual(refuse.status_code, 400)
        self.assertEqual(PresenceEtudiant.objects.get(etudiant=self.awa).statut, "PRESENT")
        accepte = self.envoyer(classeur, confirmer="true", remplacer="true")
        self.assertEqual(accepte.status_code, 201, accepte.content)
        self.assertEqual(PresenceEtudiant.objects.get(etudiant=self.awa).statut, "ABSENT")

    def test_reimport_identique_ne_cree_pas_de_doublons(self):
        classeur = self.remplir(self.modele(), {"E1": "A", "E2": "P", "E3": "P"})
        self.envoyer(classeur, confirmer="true")
        classeur = self.remplir(self.modele(), {"E1": "A", "E2": "P", "E3": "P"})
        r = self.envoyer(classeur, confirmer="true")
        self.assertEqual(r.status_code, 201, r.content)
        self.assertEqual(r.json()["resume"]["identiques"], 3)
        self.assertEqual((Seance.objects.count(), PresenceEtudiant.objects.count()), (1, 3))

    def test_seance_validee_bloque_l_import(self):
        seance = self.seance()
        self.tout_marquer(seance)
        self.client.post(f"{SEANCES}{seance.id}/valider/")
        classeur = self.remplir(self.modele(), {"E1": "A", "E2": "P", "E3": "P"})
        r = self.envoyer(classeur, confirmer="true", remplacer="true")
        self.assertEqual(r.status_code, 400)
        self.assertIn("validee", r.json()["detail"])


class ParametresTests(BasePresences):
    def test_valeurs_par_defaut_et_modification(self):
        valeurs = self.client.get("/api/parametres/").json()
        self.assertEqual(valeurs["whatsapp_actif"], "0")
        r = self.client.put("/api/parametres/", {"whatsapp_actif": "1"}, format="json")
        self.assertEqual(r.json()["whatsapp_actif"], "1")

    def test_cle_inconnue_et_role_non_administrateur_refuses(self):
        self.assertEqual(self.client.put("/api/parametres/", {"cle_pirate": "x"}, format="json").status_code, 400)
        directeur = client_de(utilisateur(Role.DIRECTEUR, "dir"))
        self.assertEqual(directeur.put("/api/parametres/", {"whatsapp_actif": "1"}, format="json").status_code, 403)


class PermissionsPresencesTests(BasePresences):
    def setUp(self):
        super().setUp()
        AffectationProfesseur.objects.create(professeur=self.prof, matiere=self.matiere, classe=self.classe, annee_scolaire=self.annee)
        self.compte_prof = utilisateur(Role.PROFESSEUR, "prof1", professeur=self.prof)
        self.client_prof = client_de(self.compte_prof)

    def test_comptable_n_a_pas_acces(self):
        self.assertEqual(client_de(utilisateur(Role.COMPTABLE, "compta")).get(SEANCES).status_code, 403)

    def test_professeur_ne_voit_que_ses_seances(self):
        mienne = self.seance()
        autre = self.seance(professeur=self.prof2, classe=self.autre_classe)
        ids = [s["id"] for s in self.client_prof.get(SEANCES).json()["resultats"]]
        self.assertEqual(ids, [mienne.id])
        self.assertEqual(self.client_prof.get(f"{SEANCES}{autre.id}/").status_code, 404)

    def test_professeur_cree_sa_seance_pour_une_classe_affectee_seulement(self):
        r = self.client_prof.post(SEANCES + "ouvrir/", self.cours())
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual(Seance.objects.get().professeur, self.prof)
        hors_affectation = self.client_prof.post(SEANCES + "ouvrir/", self.cours(classe=self.autre_classe.id, heure_debut="10:00", heure_fin="11:00"))
        self.assertEqual(hors_affectation.status_code, 400)

    def test_professeur_ne_peut_pas_modifier_sa_propre_presence(self):
        seance = self.seance()
        r = self.client_prof.patch(f"{SEANCES}{seance.id}/", {"statut_professeur": "EFFECTUE"})
        self.assertEqual(r.status_code, 400)
        seance.refresh_from_db()
        self.assertEqual(seance.statut_professeur, "PLANIFIE")
        # L'administration, elle, peut (et c'est trace).
        self.assertEqual(self.client.patch(f"{SEANCES}{seance.id}/", {"statut_professeur": "EFFECTUE"}).status_code, 200)
        self.assertTrue(JournalAudit.objects.filter(action="MODIFICATION_PRESENCE_PROFESSEUR").exists())

    def test_professeur_ne_peut_pas_rouvrir(self):
        seance = self.seance()
        self.tout_marquer(seance)
        self.client_prof.post(f"{SEANCES}{seance.id}/valider/")
        self.assertEqual(self.client_prof.post(f"{SEANCES}{seance.id}/rouvrir/", {"motif": "x"}).status_code, 403)

    def test_professeur_n_a_pas_acces_aux_notifications(self):
        self.assertEqual(self.client_prof.get("/api/presences/notifications/").status_code, 403)
