import io
from datetime import date, time
from decimal import Decimal as D
from unittest.mock import patch

import openpyxl
from django.db import IntegrityError, transaction
from django.test import TestCase
from rest_framework.test import APIClient

from apps.commun.constantes import Role
from apps.commun.models import JournalAudit
from apps.presences.models import Seance
from apps.professeurs.models import Professeur
from apps.referentiels.models import AnneeScolaire, Classe, Matiere, Niveau
from apps.remuneration.models import (
    CorrectionHeures, LigneReleve, PaiementProfesseur, ReleveRemuneration, StatutReleve, TarifProfesseurMatiere,
)
from apps.utilisateurs.models import Utilisateur

BASE = "/api/remuneration/"


def utilisateur(role=Role.ADMINISTRATEUR, identifiant="admin", **extra):
    return Utilisateur.objects.create_user(identifiant=identifiant, mot_de_passe="secret123", nom="Test", role=role, **extra)


def client_de(user):
    client = APIClient()
    client.force_authenticate(user)
    return client


class BaseRemuneration(TestCase):
    def setUp(self):
        self.annee = AnneeScolaire.objects.create(
            libelle="2025-2026", date_debut=date(2025, 9, 1), date_fin=date(2026, 6, 30), active=True)
        niveau = Niveau.objects.create(nom="6eme", code="6E", ordre=1)
        self.af = Classe.objects.create(nom="6AF", niveau=niveau, annee_scolaire=self.annee, capacite=30)
        self.d = Classe.objects.create(nom="7D", niveau=niveau, annee_scolaire=self.annee, capacite=30)
        self.maths = Matiere.objects.create(nom="Maths", code="MATH")
        self.fr = Matiere.objects.create(nom="Francais", code="FR")
        self.sy = Professeur.objects.create(matricule="P1", nom="Sy", prenom="Oumar", tarif_horaire=100)
        self.ba = Professeur.objects.create(matricule="P2", nom="Ba", prenom="Cheikh", tarif_horaire=120)
        self.fixe = Professeur.objects.create(matricule="P3", nom="Diop", prenom="Awa", mode_remuneration="FIXE", salaire_fixe=5000)
        self.admin = utilisateur()
        self.client = client_de(self.admin)
        self._heure = 8

    def seance(self, jour=6, fin_h=None, prof=None, statut="EFFECTUE", matiere=None, classe=None, mois=10, **extra):
        """Seance d'octobre 2025 ; chaque appel prend un creneau different pour respecter l'unicite."""
        debut = self._heure
        self._heure += 1
        defauts = dict(
            annee_scolaire=self.annee, classe=classe or self.af, matiere=matiere or self.maths, professeur=prof or self.sy,
            date=date(2025, mois, jour), heure_debut=time(debut, 0), heure_fin=time(fin_h or debut + 1, 0), statut_professeur=statut,
        )
        defauts.update(extra)
        return Seance.objects.create(**defauts)

    def generer(self, prof=None, mois="2025-10"):
        corps = {"mois": mois}
        if prof:
            corps["professeur"] = prof.id
        return self.client.post(BASE + "releves/generer/", corps, format="json")

    def releve(self, prof=None, mois=date(2025, 10, 1)):
        return ReleveRemuneration.objects.get(professeur=prof or self.sy, mois=mois)

    def url(self, releve):
        return f"{BASE}releves/{releve.id}/"


class CalculTests(BaseRemuneration):
    def test_seules_les_seances_effectuees_comptent(self):
        self.seance(fin_h=10)                                       # effectuee 2 h
        self.seance(statut="ABSENT")                                # absence : 0 h
        self.seance(statut="ANNULE")                                # annulee : 0 h
        self.seance(statut="PLANIFIE")                              # non validee : 0 h
        self.seance(statut="RETARD", fin_h=self._heure + 2, heure_reelle_debut=time(self._heure, 30))  # 1,5 h reelles
        self.assertEqual(self.generer(self.sy).status_code, 201)
        releve = self.releve()
        self.assertEqual(releve.lignes.count(), 2)
        detail = self.client.get(self.url(releve)).json()
        self.assertEqual(detail["totaux"]["heures"], "3.50")
        self.assertEqual(detail["totaux"]["brut"], "350.00")        # 3,5 h x 100 UM

    def test_remplacement_verse_au_remplacant_seulement(self):
        self.seance(statut="REMPLACE", professeur_remplacant=self.ba)
        self.generer(self.sy)
        self.generer(self.ba)
        self.assertEqual(ReleveRemuneration.objects.get(professeur=self.ba).lignes.count(), 1)
        self.assertEqual(self.client.get(self.url(self.releve(self.ba))).json()["totaux"]["brut"], "120.00")
        self.assertEqual(self.releve(self.sy).lignes.count(), 0)

    def test_tarif_par_matiere_prime_sur_le_tarif_general(self):
        TarifProfesseurMatiere.objects.create(professeur=self.sy, matiere=self.fr, tarif_horaire=150)
        self.seance(matiere=self.maths)                             # 1 h x 100
        self.seance(matiere=self.fr, fin_h=self._heure + 2)        # 2 h x 150
        self.generer(self.sy)
        self.assertEqual(self.client.get(self.url(self.releve())).json()["totaux"]["brut"], "400.00")

    def test_arrondi_au_centieme(self):
        TarifProfesseurMatiere.objects.create(professeur=self.sy, matiere=self.maths, tarif_horaire=D("33.33"))
        self.seance(fin_h=self._heure + 1, statut="RETARD", heure_reelle_debut=time(self._heure, 15))   # 0,75 h
        self.generer(self.sy)
        ligne = self.releve().lignes.get()
        self.assertEqual((ligne.heures, ligne.montant), (D("0.75"), D("25.00")))          # 24,9975 -> 25,00

    def test_limites_du_mois(self):
        self.seance(jour=30, mois=9)                                 # septembre : exclu
        self.seance(jour=1, mois=10)
        self.seance(jour=31, mois=10)
        self.seance(jour=1, mois=11)                                 # novembre : exclu
        self.generer(self.sy)
        self.assertEqual(sorted(self.releve().lignes.values_list("date", flat=True)), [date(2025, 10, 1), date(2025, 10, 31)])

    def test_salaire_fixe_independant_des_heures(self):
        self.seance(prof=self.fixe)
        self.assertEqual(self.generer(self.fixe).status_code, 201)
        detail = self.client.get(self.url(self.releve(self.fixe))).json()
        self.assertEqual((detail["totaux"]["brut"], detail["totaux"]["heures"]), ("5000.00", "1.00"))
        self.assertEqual(self.releve(self.fixe).lignes.get().montant, D("0"))

    def test_generation_pour_tous_les_professeurs_ignore_ceux_sans_heures(self):
        self.seance()
        self.seance(prof=self.ba)
        Professeur.objects.create(matricule="P9", nom="Inactif", prenom="X", statut="INACTIF", tarif_horaire=50)
        Professeur.objects.create(matricule="P8", nom="Sans", prenom="Heure", tarif_horaire=50)
        r = self.client.post(BASE + "releves/generer/", {"mois": "2025-10"}, format="json")
        self.assertEqual(r.json()["nombre"], 3)                                          # Sy, Ba et le salaire fixe
        self.assertEqual(sorted(x["professeur_nom"] for x in r.json()["releves"]), ["Ba Cheikh", "Diop Awa", "Sy Oumar"])

    def test_regeneration_ne_duplique_jamais_une_seance(self):
        self.seance()
        self.seance()
        for _ in range(3):
            self.generer(self.sy)
        self.assertEqual(LigneReleve.objects.count(), 2)
        self.assertEqual(ReleveRemuneration.objects.count(), 1)
        self.seance()                                                # nouvelle seance : reprise au recalcul
        self.generer(self.sy)
        self.assertEqual(LigneReleve.objects.count(), 3)

    def test_une_seance_ne_peut_etre_comptee_qu_une_fois(self):
        s = self.seance()
        self.generer(self.sy)
        autre = ReleveRemuneration.objects.create(professeur=self.ba, annee_scolaire=self.annee, mois=date(2025, 10, 1))
        with self.assertRaises(IntegrityError), transaction.atomic():
            LigneReleve.objects.create(releve=autre, seance=s, classe=s.classe, matiere=s.matiere, date=s.date,
                                       heures=1, tarif=100, montant=100)

    def test_mois_invalide_ou_hors_annee_scolaire(self):
        self.assertEqual(self.client.post(BASE + "releves/generer/", {"mois": "octobre"}, format="json").status_code, 400)
        self.seance(mois=10)
        self.assertEqual(self.generer(self.sy, mois="2027-01").status_code, 400)       # aucune annee scolaire


class ValidationTests(BaseRemuneration):
    def setUp(self):
        super().setUp()
        self.seance(fin_h=10)                                        # 2 h x 100
        self.generer(self.sy)
        self.r = self.releve()

    def valider(self):
        return self.client.post(self.url(self.r) + "valider/", {}, format="json")

    def ajouter(self, type_, montant, libelle="Ajustement"):
        return self.client.post(self.url(self.r) + "ajustement/", {"type": type_, "libelle": libelle, "montant": montant}, format="json")

    def test_primes_et_retenues_dans_le_net(self):
        self.assertEqual(self.ajouter("PRIME", "50").status_code, 200)
        r = self.ajouter("RETENUE", "30,50", "Absence reunion")
        totaux = r.json()["totaux"]
        self.assertEqual((totaux["brut"], totaux["primes"], totaux["retenues"], totaux["net"]), ("200.00", "50.00", "30.50", "219.50"))
        for corps in ({"type": "BONUS", "libelle": "x", "montant": "5"}, {"type": "PRIME", "libelle": "", "montant": "5"},
                      {"type": "PRIME", "libelle": "x", "montant": "-5"}, {"type": "PRIME", "libelle": "x", "montant": "1.234"}):
            self.assertEqual(self.client.post(self.url(self.r) + "ajustement/", corps, format="json").status_code, 400, corps)

    def test_suppression_d_un_ajustement(self):
        ajustement = self.ajouter("PRIME", "50").json()["ajustements"][0]
        r = self.client.post(self.url(self.r) + "supprimer-ajustement/", {"ajustement": ajustement["id"]}, format="json")
        self.assertEqual(r.json()["totaux"]["net"], "200.00")

    def test_validation_fige_les_montants(self):
        self.ajouter("PRIME", "50")
        detail = self.valider().json()
        self.assertEqual((detail["statut"], detail["totaux"]["net"]), ("VALIDE", "250.00"))
        # Un changement de tarif ulterieur ne modifie pas un releve valide.
        self.sy.tarif_horaire = 999
        self.sy.save()
        self.assertEqual(self.client.get(self.url(self.r)).json()["totaux"]["net"], "250.00")
        self.assertEqual(self.ajouter("PRIME", "10").status_code, 400)                  # ajustements fermes
        self.assertEqual(self.generer(self.sy).status_code, 400)                        # recalcul interdit
        self.assertTrue(JournalAudit.objects.filter(action="VALIDATION_RELEVE").exists())
        self.assertEqual(self.valider().status_code, 400)                                # deja valide

    def test_validation_refusee_si_vide_ou_net_negatif(self):
        self.ajouter("RETENUE", "500")
        r = self.valider()
        self.assertEqual(r.status_code, 400)
        self.assertIn("negatif", r.json()["detail"])
        LigneReleve.objects.all().delete()
        self.client.post(self.url(self.r) + "supprimer-ajustement/", {"ajustement": self.r.ajustements.get().id}, format="json")
        self.assertIn("rien a valider", self.valider().json()["detail"])

    def test_seance_comptee_dans_un_releve_valide_non_modifiable(self):
        seance = Seance.objects.get()
        url = f"/api/presences/seances/{seance.id}/"
        self.assertEqual(self.client.patch(url, {"observations_professeur": "ok"}).status_code, 200)   # brouillon : possible
        self.valider()
        for corps in ({"statut_professeur": "ABSENT"}, {"heure_fin": "11:00"}, {"date": "2025-10-07"}):
            self.assertEqual(self.client.patch(url, corps).status_code, 400, corps)
        self.assertEqual(self.client.delete(url).status_code, 400)
        seance.refresh_from_db()
        self.assertEqual((seance.statut_professeur, seance.heures_effectuees), ("EFFECTUE", D("2.00")))

    def test_reouverture_avec_motif_et_sans_paiement(self):
        self.valider()
        self.assertEqual(self.client.post(self.url(self.r) + "rouvrir/", {}, format="json").status_code, 400)
        compta = client_de(utilisateur(Role.COMPTABLE, "compta"))
        self.assertEqual(compta.post(self.url(self.r) + "rouvrir/", {"motif": "x"}, format="json").status_code, 403)
        self.client.post(self.url(self.r) + "paiement/", {"montant": "50"}, format="json")
        self.assertEqual(self.client.post(self.url(self.r) + "rouvrir/", {"motif": "x"}, format="json").status_code, 400)  # paiements
        PaiementProfesseur.objects.all()  # (immuables : ils restent)
        autre = self.seance(prof=self.ba)
        self.generer(self.ba)
        self.client.post(self.url(self.releve(self.ba)) + "valider/", {}, format="json")
        r = self.client.post(self.url(self.releve(self.ba)) + "rouvrir/", {"motif": "oubli d'une prime"}, format="json")
        self.assertEqual((r.status_code, r.json()["statut"]), (200, "BROUILLON"))
        self.assertTrue(JournalAudit.objects.filter(action="REOUVERTURE_RELEVE").exists())
        self.assertIsNotNone(autre)

    def test_validation_reservee_a_la_direction(self):
        compta = client_de(utilisateur(Role.COMPTABLE, "compta"))
        self.assertEqual(compta.post(self.url(self.r) + "valider/", {}, format="json").status_code, 403)
        # Le comptable peut preparer le releve (generation, primes) mais pas le valider.
        self.assertEqual(compta.post(self.url(self.r) + "ajustement/", {"type": "PRIME", "libelle": "x", "montant": "5"}, format="json").status_code, 200)


class PaiementTests(BaseRemuneration):
    def setUp(self):
        super().setUp()
        self.seance(fin_h=11)                                        # 3 h x 100 = 300
        self.generer(self.sy)
        self.r = self.releve()
        self.paiement_url = self.url(self.r) + "paiement/"

    def valider(self):
        return self.client.post(self.url(self.r) + "valider/", {}, format="json")

    def payer(self, montant, **extra):
        return self.client.post(self.paiement_url, {"montant": montant, "date_paiement": "2025-11-05", **extra}, format="json")

    def test_paiement_refuse_avant_validation_des_heures(self):
        self.assertEqual(self.payer("100").status_code, 400)

    def test_acomptes_partiels_solde_et_statuts(self):
        self.valider()
        r = self.payer("100").json()
        self.assertEqual((r["totaux"]["paye"], r["totaux"]["solde"], r["totaux"]["statut_paiement"]), ("100.00", "200.00", "PARTIEL"))
        self.assertEqual(r["paiements"][0]["type"], "ACOMPTE")
        r = self.payer("200").json()
        self.assertEqual((r["totaux"]["solde"], r["totaux"]["statut_paiement"]), ("0.00", "PAYE"))
        self.assertEqual({p["type"] for p in r["paiements"]}, {"ACOMPTE", "SOLDE"})
        self.assertTrue(all(p["reference"].startswith("REM-") for p in r["paiements"]))
        self.assertEqual(self.payer("1").status_code, 400)                              # plus rien a payer

    def test_paiement_superieur_au_solde_ou_invalide_refuse(self):
        self.valider()
        self.payer("250")
        for montant in ("51", "0", "-10", "abc", "10.555"):
            self.assertEqual(self.payer(montant).status_code, 400, montant)
        self.assertEqual(self.payer("10", mode_paiement="TROC").status_code, 400)
        self.assertEqual(PaiementProfesseur.objects.count(), 1)

    def test_paiements_immuables(self):
        self.valider()
        self.payer("100")
        paiement = PaiementProfesseur.objects.get()
        paiement.montant = 1
        with self.assertRaises(TypeError):
            paiement.save()
        with self.assertRaises(TypeError):
            paiement.delete()
        with self.assertRaises(TypeError):
            PaiementProfesseur.objects.all().update(montant=0)

    def test_annulation_par_paiement_inverse(self):
        self.valider()
        self.payer("100")
        paiement = PaiementProfesseur.objects.get()
        url = f"{BASE}paiements/{paiement.id}/annuler/"
        self.assertEqual(self.client.post(url, {}, format="json").status_code, 400)                      # motif obligatoire
        compta = client_de(utilisateur(Role.COMPTABLE, "compta"))
        self.assertEqual(compta.post(url, {"motif": "erreur"}, format="json").status_code, 403)          # direction seulement
        r = self.client.post(url, {"motif": "montant saisi par erreur"}, format="json")
        self.assertEqual(r.status_code, 201, r.content)
        self.assertEqual((r.json()["type"], r.json()["montant"]), ("ANNULATION", "-100.00"))
        paiement.refresh_from_db()
        self.assertEqual(paiement.montant, D("100"))                                                    # l'original n'est pas modifie
        detail = self.client.get(self.url(self.r)).json()
        self.assertEqual((detail["totaux"]["paye"], detail["totaux"]["statut_paiement"]), ("0.00", "NON_PAYE"))
        self.assertEqual(self.client.post(url, {"motif": "encore"}, format="json").status_code, 400)    # deja annule
        annulation = PaiementProfesseur.objects.get(type="ANNULATION")
        self.assertEqual(self.client.post(f"{BASE}paiements/{annulation.id}/annuler/", {"motif": "x"}, format="json").status_code, 400)
        self.assertTrue(JournalAudit.objects.filter(action="ANNULATION_PAIEMENT").exists())
        self.assertEqual(self.payer("300").status_code, 200)                                            # on peut repayer le tout
        paiements = self.client.get(BASE + "paiements/", {"releve": self.r.id}).json()["resultats"]
        self.assertEqual(len(paiements), 3)
        self.assertEqual([p["annulable"] for p in paiements].count(True), 1)

    def test_historique_des_paiements_filtrable(self):
        self.valider()
        self.payer("100")
        self.assertEqual(self.client.get(BASE + "paiements/", {"professeur": self.sy.id}).json()["total"], 1)
        self.assertEqual(self.client.get(BASE + "paiements/", {"professeur": self.ba.id}).json()["total"], 0)
        self.assertEqual(self.client.get(BASE + "paiements/", {"du": "2025-12-01"}).json()["total"], 0)
        self.assertEqual(self.client.get(BASE + "paiements/", {"mois": "2025-10"}).json()["total"], 1)


class CorrectionHeuresTests(BaseRemuneration):
    def setUp(self):
        super().setUp()
        self.seance(fin_h=10)
        self.seance(fin_h=self._heure + 1)
        self.generer(self.sy)
        self.r = self.releve()
        self.client.post(self.url(self.r) + "valider/", {}, format="json")        # 3 h : 300
        self.ligne = self.r.lignes.get(heures=2)                                    # la seance de 2 h

    def corriger(self, heures, motif="Cours ecourte", client=None):
        return (client or self.client).post(self.url(self.r) + "corriger-heures/",
                                            {"ligne": self.ligne.id, "heures": heures, "motif": motif}, format="json")

    def test_correction_motivee_recalcule_et_trace(self):
        r = self.corriger("1,5")
        self.assertEqual(r.status_code, 200, r.content)
        totaux = r.json()["totaux"]
        self.assertEqual((totaux["heures"], totaux["brut"], totaux["net"]), ("2.50", "250.00", "250.00"))
        correction = CorrectionHeures.objects.get()
        self.assertEqual((correction.anciennes_heures, correction.nouvelles_heures, correction.ancien_montant, correction.nouveau_montant),
                         (D("2"), D("1.5"), D("200"), D("150")))
        self.assertEqual(correction.motif, "Cours ecourte")
        self.assertEqual(r.json()["corrections"][0]["motif"], "Cours ecourte")
        audit = JournalAudit.objects.get(action="CORRECTION_HEURES")
        self.assertEqual(audit.nouvelle_valeur["motif"], "Cours ecourte")

    def test_correction_exige_motif_direction_et_valeur_valide(self):
        self.assertEqual(self.corriger("1", motif="").status_code, 400)
        compta = client_de(utilisateur(Role.COMPTABLE, "compta"))
        self.assertEqual(self.corriger("1", client=compta).status_code, 403)
        for valeur in ("-1", "abc", "1.234", "2"):                                  # "2" = inchangee
            self.assertEqual(self.corriger(valeur).status_code, 400, valeur)
        self.assertEqual(CorrectionHeures.objects.count(), 0)

    def test_correction_refusee_si_le_net_passe_sous_les_paiements(self):
        self.client.post(self.url(self.r) + "paiement/", {"montant": "300"}, format="json")
        r = self.corriger("0")                                                      # net 100 < 300 deja payes
        self.assertEqual(r.status_code, 400)
        self.assertIn("paiements", r.json()["detail"])
        self.assertEqual(self.corriger("3").status_code, 200)                       # hausse : possible, statut PARTIEL
        self.assertEqual(self.client.get(self.url(self.r)).json()["totaux"]["statut_paiement"], "PARTIEL")

    def test_ligne_d_un_autre_releve_ou_releve_brouillon(self):
        autre = self.seance(prof=self.ba)
        self.generer(self.ba)
        ligne_ba = LigneReleve.objects.get(seance=autre)
        r = self.client.post(self.url(self.r) + "corriger-heures/", {"ligne": ligne_ba.id, "heures": "1", "motif": "x"}, format="json")
        self.assertEqual(r.status_code, 404)
        brouillon = self.releve(self.ba)
        r = self.client.post(self.url(brouillon) + "corriger-heures/", {"ligne": ligne_ba.id, "heures": "0.5", "motif": "x"}, format="json")
        self.assertEqual(r.status_code, 400)


class ConsultationTests(BaseRemuneration):
    def test_heures_par_matiere_et_par_classe(self):
        self.seance(matiere=self.maths, classe=self.af, fin_h=10)
        self.seance(matiere=self.maths, classe=self.d)
        self.seance(matiere=self.fr, classe=self.af)
        self.seance(statut="ABSENT")
        self.generer(self.sy)
        self.seance(matiere=self.maths, classe=self.af, jour=20)                    # nouvelle seance, pas encore comptabilisee
        r = self.client.get(BASE + "releves/heures/", {"professeur": self.sy.id, "du": "2025-10-01", "au": "2025-10-31"}).json()
        self.assertEqual(r["total_heures"], "5.00")                                  # 2 + 1 + 1 + 1 (l'absence ne compte pas)
        lignes = {(l["matiere_nom"], l["classe_nom"]): l for l in r["detail"]}
        self.assertEqual(lignes[("Maths", "6AF")]["heures"], "3.00")
        self.assertEqual(lignes[("Maths", "6AF")]["heures_comptabilisees"], "2.00")
        self.assertEqual(lignes[("Maths", "6AF")]["heures_a_comptabiliser"], "1.00")
        self.assertEqual(lignes[("Maths", "7D")]["heures_a_comptabiliser"], "0.00")
        self.assertEqual(client_de(utilisateur(Role.COMPTABLE, "compta")).get(BASE + "releves/heures/", {"professeur": self.sy.id, "du": "x", "au": "y"}).status_code, 400)

    def test_etat_mensuel_totaux_et_professeurs_sans_releve(self):
        self.seance(fin_h=10)
        self.seance(prof=self.ba)
        self.generer(self.sy)
        etat = self.client.get(BASE + "releves/etat/", {"mois": "2025-10"}).json()
        self.assertEqual([r["professeur_nom"] for r in etat["releves"]], ["Sy Oumar"])
        self.assertEqual([s["nom"] for s in etat["sans_releve"]], ["Ba Cheikh"])
        self.assertEqual((etat["total"]["heures"], etat["total"]["net"]), ("2.00", "200.00"))          # brouillon : calcul courant
        self.assertEqual(etat["total_valide"]["net"], "0.00")                                          # ... mais rien de valide
        self.client.post(self.url(self.releve()) + "valider/", {}, format="json")
        self.client.post(self.url(self.releve()) + "paiement/", {"montant": "50"}, format="json")
        etat = self.client.get(BASE + "releves/etat/", {"mois": "2025-10"}).json()
        self.assertEqual((etat["total"]["net"], etat["total"]["paye"], etat["total"]["solde"]), ("200.00", "50.00", "150.00"))
        self.assertEqual(etat["total_valide"], etat["total"])

    def test_filtres_de_la_liste(self):
        self.seance(matiere=self.maths)
        self.seance(prof=self.ba, matiere=self.fr)
        self.generer(self.sy)
        self.generer(self.ba)
        for prof in (self.sy, self.ba):
            self.client.post(self.url(self.releve(prof)) + "valider/", {}, format="json")
        self.client.post(self.url(self.releve(self.ba)) + "paiement/", {"montant": "120"}, format="json")
        noms = lambda **p: [r["professeur_nom"] for r in self.client.get(BASE + "releves/", p).json()["resultats"]]
        self.assertEqual(sorted(noms()), ["Ba Cheikh", "Sy Oumar"])
        self.assertEqual(noms(professeur=self.sy.id), ["Sy Oumar"])
        self.assertEqual(noms(statut_paiement="PAYE"), ["Ba Cheikh"])
        self.assertEqual(noms(statut_paiement="NON_PAYE"), ["Sy Oumar"])
        self.assertEqual(noms(matiere=self.fr.id), ["Ba Cheikh"])
        self.assertEqual(noms(mois="2025-11"), [])
        self.assertEqual(noms(statut="BROUILLON"), [])


class DocumentsTests(BaseRemuneration):
    def setUp(self):
        super().setUp()
        self.seance(fin_h=10)
        self.generer(self.sy)
        self.r = self.releve()

    def pdf(self, url, **params):
        with patch("reportlab.rl_config.pageCompression", 0):
            return self.client.get(url, params)

    def test_fiche_provisoire_puis_valide(self):
        provisoire = self.pdf(self.url(self.r) + "pdf/")
        self.assertEqual(provisoire["Content-Type"], "application/pdf")
        self.assertTrue(provisoire.content.startswith(b"%PDF"))
        self.assertIn(b"PROVISOIRE", provisoire.content)
        self.client.post(self.url(self.r) + "valider/", {}, format="json")
        self.assertNotIn(b"PROVISOIRE", self.pdf(self.url(self.r) + "pdf/").content)

    def test_libelle_avec_caracteres_speciaux_ne_casse_pas_le_pdf(self):
        self.client.post(self.url(self.r) + "ajustement/", {"type": "PRIME", "libelle": "Prime <b>assiduite</b> & rendement", "montant": "40"}, format="json")
        self.assertEqual(self.pdf(self.url(self.r) + "pdf/").status_code, 200)

    def test_etat_pdf_et_excel(self):
        self.client.post(self.url(self.r) + "valider/", {}, format="json")
        self.assertTrue(self.pdf(BASE + "releves/etat/pdf/", mois="2025-10").content.startswith(b"%PDF"))
        classeur = openpyxl.load_workbook(io.BytesIO(self.client.get(BASE + "releves/etat/excel/", {"mois": "2025-10"}).content))
        lignes = [[c.value for c in ligne] for ligne in classeur["Etat de paiement"].iter_rows()]
        sy = next(l for l in lignes if l[0] == "Sy Oumar")
        self.assertEqual((sy[2], sy[3], sy[6]), (2.0, 200.0, 200.0))           # heures, brut, net
        total = next(l for l in lignes if l[0] == "TOTAL")
        self.assertEqual(total[6], 200.0)
        self.assertEqual(len(list(classeur["Detail des seances"].iter_rows())), 2)
        self.assertEqual(self.client.get(BASE + "releves/etat/pdf/", {"mois": "n'importe quoi"}).status_code, 400)


class PermissionsRemunerationTests(BaseRemuneration):
    def setUp(self):
        super().setUp()
        self.seance()
        self.seance(prof=self.ba)
        self.generer(self.sy)
        self.generer(self.ba)
        self.compte_sy = utilisateur(Role.PROFESSEUR, "prof_sy", professeur=self.sy)
        self.prof = client_de(self.compte_sy)

    def test_professeur_ne_voit_que_ses_releves_valides(self):
        self.assertEqual(self.prof.get(BASE + "releves/").json()["total"], 0)               # brouillons invisibles
        for p in (self.sy, self.ba):
            self.client.post(self.url(self.releve(p)) + "valider/", {}, format="json")
        liste = self.prof.get(BASE + "releves/").json()
        self.assertEqual([r["professeur_nom"] for r in liste["resultats"]], ["Sy Oumar"])
        self.assertEqual(self.prof.get(self.url(self.releve(self.sy))).status_code, 200)
        self.assertEqual(self.prof.get(self.url(self.releve(self.ba))).status_code, 404)
        self.assertEqual(self.prof.get(self.url(self.releve(self.sy)) + "pdf/").status_code, 200)
        self.assertEqual(self.prof.get(self.url(self.releve(self.ba)) + "pdf/").status_code, 404)

    def test_professeur_ne_peut_rien_ecrire(self):
        self.client.post(self.url(self.releve(self.sy)) + "valider/", {}, format="json")
        url = self.url(self.releve(self.sy))
        self.assertEqual(self.prof.post(BASE + "releves/generer/", {"mois": "2025-10"}, format="json").status_code, 403)
        self.assertEqual(self.prof.post(url + "paiement/", {"montant": "10"}, format="json").status_code, 403)
        self.assertEqual(self.prof.post(url + "ajustement/", {"type": "PRIME", "libelle": "x", "montant": "10"}, format="json").status_code, 403)
        self.assertEqual(self.prof.get(BASE + "paiements/").status_code, 403)
        self.assertEqual(self.prof.get(BASE + "tarifs-matiere/").status_code, 403)
        self.assertEqual(self.prof.get(BASE + "releves/etat/", {"mois": "2025-10"}).status_code, 403)

    def test_professeur_consulte_uniquement_ses_propres_heures(self):
        r = self.prof.get(BASE + "releves/heures/", {"professeur": self.ba.id, "du": "2025-10-01", "au": "2025-10-31"})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.json()["total_heures"], "1.00")                                  # ses propres heures (Sy), pas celles de Ba

    def test_autres_roles(self):
        self.assertEqual(client_de(utilisateur(Role.RESPONSABLE_SCOLARITE, "scol")).get(BASE + "releves/").status_code, 403)
        self.assertEqual(client_de(utilisateur(Role.RESPONSABLE_PEDAGOGIQUE, "peda")).get(BASE + "releves/").status_code, 403)
        compta = client_de(utilisateur(Role.COMPTABLE, "compta"))
        self.assertEqual(compta.get(BASE + "releves/").json()["total"], 2)
        self.assertEqual(compta.post(BASE + "tarifs-matiere/", {"professeur": self.sy.id, "matiere": self.fr.id, "tarif_horaire": "140"}).status_code, 201)
        self.assertEqual(compta.post(BASE + "tarifs-matiere/", {"professeur": self.sy.id, "matiere": self.fr.id, "tarif_horaire": "150"}).status_code, 400)  # doublon
        self.assertEqual(compta.post(BASE + "tarifs-matiere/", {"professeur": self.sy.id, "matiere": self.maths.id, "tarif_horaire": "0"}).status_code, 400)
