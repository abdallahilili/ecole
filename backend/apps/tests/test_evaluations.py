import io
from datetime import date
from decimal import Decimal as D
from unittest.mock import patch

import openpyxl
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from rest_framework.test import APIClient

from apps.commun.constantes import Role
from apps.commun.models import JournalAudit, Parametre
from apps.etudiants.models import Etudiant, Inscription
from apps.evaluations.models import Appreciation, Evaluation, Note, StatutNote, ValidationResultats
from apps.professeurs.models import AffectationProfesseur, Professeur
from apps.referentiels.models import AffectationMatiereClasse, AnneeScolaire, Classe, Matiere, Niveau
from apps.utilisateurs.models import Utilisateur

EVALS = "/api/evaluations/evaluations/"
RESULTATS = "/api/evaluations/resultats/"


def utilisateur(role=Role.ADMINISTRATEUR, identifiant="admin", **extra):
    return Utilisateur.objects.create_user(identifiant=identifiant, mot_de_passe="secret123", nom="Test", role=role, **extra)


def client_de(user):
    client = APIClient()
    client.force_authenticate(user)
    return client


class BaseEvaluations(TestCase):
    def setUp(self):
        self.annee = AnneeScolaire.objects.create(
            libelle="2025-2026", date_debut=date(2025, 9, 1), date_fin=date(2026, 6, 30), active=True)
        self.niveau = Niveau.objects.create(nom="1AS", code="1AS", ordre=1)
        self.classe = Classe.objects.create(nom="1As", niveau=self.niveau, annee_scolaire=self.annee, capacite=40)
        self.maths = Matiere.objects.create(nom="Maths", code="MATH")
        self.fr = Matiere.objects.create(nom="Francais", code="FR")
        self.sans_coef = Matiere.objects.create(nom="Dessin", code="DESS")
        AffectationMatiereClasse.objects.create(matiere=self.maths, classe=self.classe, coefficient=5)
        AffectationMatiereClasse.objects.create(matiere=self.fr, classe=self.classe, coefficient=3)
        self.awa = self.eleve("E1", "Diallo", "Awa")
        self.moussa = self.eleve("E2", "Ba", "Moussa")
        self.fall = self.eleve("E3", "Fall", "Mariam")
        self.admin = utilisateur()
        self.client = client_de(self.admin)

    def eleve(self, matricule, nom, prenom, classe=None):
        etudiant = Etudiant.objects.create(matricule=matricule, nom=nom, prenom=prenom)
        Inscription.objects.create(etudiant=etudiant, annee_scolaire=self.annee, classe=classe or self.classe,
                                   date_inscription=date(2025, 9, 1))
        return etudiant

    def evaluation(self, matiere, libelle, poids, trimestre=1, bareme=20, type="EXAMEN"):
        return Evaluation.objects.create(
            annee_scolaire=self.annee, classe=self.classe, matiere=matiere, trimestre=trimestre,
            libelle=libelle, type=type, bareme=bareme, poids=poids)

    def noter(self, evaluation, etudiant, valeur, statut=StatutNote.SAISIE):
        return Note.objects.create(evaluation=evaluation, etudiant=etudiant, statut=statut,
                                   valeur=None if valeur is None else D(str(valeur)))

    def jeu_complet(self, trimestre=1):
        """Interrogations (x3), composition 1 (x1), composition 2 (x2) en Maths et Francais."""
        evals = {}
        for matiere in (self.maths, self.fr):
            evals[matiere.code] = [self.evaluation(matiere, "Interros", 3, trimestre, type="INTERROGATION"),
                                   self.evaluation(matiere, "Composition 1", 1, trimestre),
                                   self.evaluation(matiere, "Composition 2", 2, trimestre)]
        donnees = {
            "MATH": {self.awa: (4.75, 3, 10), self.moussa: (16, 14, 12), self.fall: (10, None, None)},
            "FR": {self.awa: (12, 6, 9.5), self.moussa: (15, 12, 10)},
        }
        for code, par_eleve in donnees.items():
            for etudiant, valeurs in par_eleve.items():
                for evaluation, valeur in zip(evals[code], valeurs):
                    if valeur is not None:
                        self.noter(evaluation, etudiant, valeur)
        return evals

    def resultats(self, periode=1, **extra):
        reponse = self.client.get(RESULTATS, {"classe": self.classe.id, "periode": periode, **extra})
        self.assertEqual(reponse.status_code, 200, reponse.content)
        return reponse.json()

    @staticmethod
    def ligne(donnees, etudiant):
        return next(l for l in donnees["etudiants"] if l["etudiant"] == etudiant.id)


class EvaluationApiTests(BaseEvaluations):
    def corps(self, **extra):
        return {"annee_scolaire": self.annee.id, "classe": self.classe.id, "matiere": self.maths.id,
                "trimestre": 1, "libelle": "Devoir 1", "type": "DEVOIR", "bareme": "20", "poids": "1", **extra}

    def test_creation_et_regles(self):
        self.assertEqual(self.client.post(EVALS, self.corps()).status_code, 201)
        self.assertEqual(self.client.post(EVALS, self.corps()).status_code, 400)                        # libelle en double
        self.assertEqual(self.client.post(EVALS, self.corps(libelle="B", bareme="0")).status_code, 400)
        self.assertEqual(self.client.post(EVALS, self.corps(libelle="C", poids="-1")).status_code, 400)
        self.assertEqual(self.client.post(EVALS, self.corps(libelle="D", date="2027-01-01")).status_code, 400)
        r = self.client.post(EVALS, self.corps(libelle="E", matiere=self.sans_coef.id))
        self.assertEqual(r.status_code, 400)
        self.assertIn("coefficient", r.json()["detail"])

    def test_bareme_fige_des_que_des_notes_existent(self):
        evaluation = self.evaluation(self.maths, "Devoir", 1)
        self.assertEqual(self.client.patch(f"{EVALS}{evaluation.id}/", {"bareme": "10"}).status_code, 200)
        self.noter(evaluation, self.awa, 8)
        self.assertEqual(self.client.patch(f"{EVALS}{evaluation.id}/", {"bareme": "20"}).status_code, 400)
        self.assertEqual(self.client.patch(f"{EVALS}{evaluation.id}/", {"poids": "2"}).status_code, 200)
        self.assertEqual(self.client.delete(f"{EVALS}{evaluation.id}/").status_code, 400)

    def test_suppression_d_une_evaluation_vide(self):
        evaluation = self.evaluation(self.maths, "Devoir", 1)
        self.assertEqual(self.client.delete(f"{EVALS}{evaluation.id}/").status_code, 204)


class SaisieNotesTests(BaseEvaluations):
    def setUp(self):
        super().setUp()
        self.evaluation_ = self.evaluation(self.maths, "Devoir", 1, bareme=20)
        self.url = f"{EVALS}{self.evaluation_.id}/notes/"

    def envoyer(self, lignes, **extra):
        return self.client.post(self.url, {"notes": lignes, **extra}, format="json")

    def test_liste_marque_les_notes_manquantes(self):
        self.noter(self.evaluation_, self.awa, 14)
        eleves = {e["matricule"]: e for e in self.client.get(self.url).json()["eleves"]}
        self.assertEqual((eleves["E1"]["statut"], eleves["E1"]["valeur"]), ("SAISIE", "14"))
        self.assertEqual(eleves["E2"]["statut"], "MANQUANTE")
        self.assertIsNone(eleves["E2"]["valeur"])

    def test_saisie_virgule_decimale_et_statuts(self):
        r = self.envoyer([
            {"etudiant": self.awa.id, "statut": "SAISIE", "valeur": "14,5"},
            {"etudiant": self.moussa.id, "statut": "ABSENT_JUSTIFIE"},
            {"etudiant": self.fall.id, "statut": "NON_APPLICABLE", "valeur": "12"},
        ])
        self.assertEqual(r.status_code, 200, r.content)
        notes = {n.etudiant.matricule: (n.statut, n.valeur) for n in Note.objects.all()}
        self.assertEqual(notes, {"E1": ("SAISIE", D("14.50")), "E2": ("ABSENT_JUSTIFIE", None), "E3": ("NON_APPLICABLE", None)})

    def test_note_invalide_rien_n_est_enregistre(self):
        for valeur in ("20,5", "-1", "abc", "12.345", ""):
            r = self.envoyer([
                {"etudiant": self.moussa.id, "statut": "SAISIE", "valeur": "10"},
                {"etudiant": self.awa.id, "statut": "SAISIE", "valeur": valeur},
            ])
            self.assertEqual(r.status_code, 400, valeur)
            self.assertIn("Diallo Awa", r.json()["detail"])
        self.assertEqual(Note.objects.count(), 0)

    def test_eleve_hors_classe_et_doublon_refuses(self):
        hors = self.eleve("E9", "Sow", "Ali", classe=Classe.objects.create(
            nom="2As", niveau=self.niveau, annee_scolaire=self.annee, capacite=30))
        self.assertEqual(self.envoyer([{"etudiant": hors.id, "statut": "SAISIE", "valeur": "10"}]).status_code, 400)
        double = [{"etudiant": self.awa.id, "statut": "SAISIE", "valeur": "10"}] * 2
        self.assertEqual(self.envoyer(double).status_code, 400)

    def test_une_note_n_est_jamais_transformee_en_zero_et_effacement_audite(self):
        self.envoyer([{"etudiant": self.awa.id, "statut": "SAISIE", "valeur": "15"}])
        self.assertEqual(self.envoyer([{"etudiant": self.awa.id, "statut": "SAISIE", "valeur": "16"}]).status_code, 200)
        audit = JournalAudit.objects.get(action="MODIFICATION_NOTE")
        self.assertEqual((audit.ancienne_valeur["valeur"], audit.nouvelle_valeur["valeur"]), ("15.00", "16.00"))
        self.envoyer([{"etudiant": self.awa.id, "statut": "MANQUANTE"}])
        self.assertEqual(Note.objects.count(), 0)
        self.assertTrue(JournalAudit.objects.filter(action="SUPPRESSION_NOTE").exists())


class ResultatsTests(BaseEvaluations):
    def test_moyennes_moyenne_generale_et_rang(self):
        self.jeu_complet()
        donnees = self.resultats()
        awa, moussa, fall = (self.ligne(donnees, e) for e in (self.awa, self.moussa, self.fall))
        # Awa : maths (3x4,75+3+2x10)/6 = 6,2083 ; francais (3x12+6+2x9,5)/6 = 10,1667 ; MG = (5x6,2083+3x10,1667)/8
        self.assertEqual(awa["matieres"][str(self.maths.id)]["moyenne"], "6.21")
        self.assertEqual(awa["matieres"][str(self.fr.id)]["moyenne"], "10.17")
        self.assertEqual(awa["moyenne_generale"], "7.69")
        self.assertEqual(moussa["moyenne_generale"], "13.77")
        self.assertEqual((moussa["rang"], fall["rang"], awa["rang"]), (1, 2, 3))
        self.assertEqual(donnees["effectif"], 3)
        self.assertEqual(donnees["statistiques"]["plus_haute"], "13.77")

    def test_notes_manquantes_signalees_jamais_zero(self):
        self.jeu_complet()
        fall = self.ligne(self.resultats(), self.fall)
        self.assertTrue(fall["incomplet"])
        self.assertEqual(fall["matieres"][str(self.maths.id)]["moyenne"], "10.00")     # 2 evaluations manquantes ignorees
        self.assertTrue(fall["matieres"][str(self.maths.id)]["incomplete"])
        self.assertIsNone(fall["matieres"][str(self.fr.id)]["moyenne"])
        self.assertEqual(fall["matieres_manquantes"], ["Francais"])
        self.assertEqual(fall["moyenne_generale"], "10.00")

    def test_regle_matiere_non_evaluee_comptee_zero_comme_le_papier(self):
        self.jeu_complet()
        Parametre.objects.create(cle="matiere_non_evaluee", valeur="zero")
        donnees = self.resultats()
        self.assertEqual(self.ligne(donnees, self.fall)["moyenne_generale"], "6.25")     # 50 / 8
        self.assertEqual(self.ligne(donnees, self.fall)["rang"], 3)

    def test_eleve_sans_aucune_note_n_a_pas_de_rang(self):
        self.jeu_complet()
        sans_note = self.eleve("E4", "Sy", "Omar")
        ligne = self.ligne(self.resultats(), sans_note)
        self.assertIsNone(ligne["moyenne_generale"])
        self.assertIsNone(ligne["rang"])

    def test_departage_alphabetique_configurable(self):
        for etudiant in (self.awa, self.moussa):
            for matiere in (self.maths, self.fr):
                evaluation = Evaluation.objects.get_or_create(
                    annee_scolaire=self.annee, classe=self.classe, matiere=matiere, trimestre=1, libelle="Devoir",
                    defaults={"poids": 1})[0]
                self.noter(evaluation, etudiant, 14)
        donnees = self.resultats()
        self.assertEqual((self.ligne(donnees, self.awa)["rang"], self.ligne(donnees, self.moussa)["rang"]), (1, 1))
        Parametre.objects.create(cle="ex_aequo", valeur="alphabetique")
        donnees = self.resultats()
        self.assertEqual((self.ligne(donnees, self.moussa)["rang"], self.ligne(donnees, self.awa)["rang"]), (1, 2))  # Ba < Diallo

    def test_resultat_annuel_pondere_par_trimestre(self):
        for trimestre, valeur in ((1, 10), (2, 12), (3, 16)):
            for matiere in (self.maths, self.fr):
                self.noter(self.evaluation(matiere, "Devoir", 1, trimestre), self.awa, valeur)
        Parametre.objects.create(cle="poids_trimestres", valeur="1,1,2")
        annuel = self.resultats(periode=0)
        awa = self.ligne(annuel, self.awa)
        self.assertEqual(awa["moyennes_trimestres"], {"1": "10.00", "2": "12.00", "3": "16.00"})
        self.assertEqual(awa["moyenne_generale"], "13.50")
        self.assertFalse(awa["incomplet"])
        self.assertFalse(annuel["validee"])

    def test_periode_ou_classe_invalide(self):
        self.assertEqual(self.client.get(RESULTATS, {"classe": self.classe.id, "periode": 7}).status_code, 400)
        self.assertEqual(self.client.get(RESULTATS).status_code, 400)


class ValidationEtVerrouTests(BaseEvaluations):
    def setUp(self):
        super().setUp()
        self.evals = self.jeu_complet()
        self.valider = lambda **extra: self.client.post(RESULTATS + "valider/", {"classe": self.classe.id, "periode": 1, **extra}, format="json")
        self.notes_url = f"{EVALS}{self.evals['MATH'][0].id}/notes/"

    def modifier_note(self, client, **extra):
        return client.post(self.notes_url, {"notes": [{"etudiant": self.awa.id, "statut": "SAISIE", "valeur": "18"}], **extra}, format="json")

    def test_validation_exige_la_completude_sauf_acceptation_explicite(self):
        r = self.valider()
        self.assertEqual(r.status_code, 400)
        self.assertIn("manquantes", r.json()["detail"])
        self.assertFalse(ValidationResultats.objects.exists())
        self.assertEqual(self.valider(accepter_incomplet=True).status_code, 200)
        self.assertTrue(JournalAudit.objects.filter(action="VALIDATION_RESULTATS").exists())
        self.assertEqual(self.valider(accepter_incomplet=True).status_code, 400)       # deja valides

    def test_apres_validation_notes_verrouillees_correction_tracee_et_motivee(self):
        self.valider(accepter_incomplet=True)
        self.assertTrue(self.resultats()["validee"])
        # Un professeur affecte ne peut plus rien modifier.
        prof = Professeur.objects.create(matricule="P1", nom="Sy", prenom="Oumar", tarif_horaire=100)
        AffectationProfesseur.objects.create(professeur=prof, matiere=self.maths, classe=self.classe, annee_scolaire=self.annee)
        client_prof = client_de(utilisateur(Role.PROFESSEUR, "prof", professeur=prof))
        r = self.modifier_note(client_prof, motif="pitie")
        self.assertEqual(r.status_code, 400)
        # La direction doit justifier.
        self.assertEqual(self.modifier_note(self.client).status_code, 400)
        r = self.modifier_note(self.client, motif="erreur de recopie")
        self.assertEqual(r.status_code, 200, r.content)
        audit = JournalAudit.objects.filter(action="MODIFICATION_NOTE").latest("id")
        self.assertEqual(audit.nouvelle_valeur["motif"], "erreur de recopie")
        self.assertEqual(Note.objects.get(evaluation=self.evals["MATH"][0], etudiant=self.awa).valeur, D("18"))
        # Evaluations figees elles aussi.
        corps = {"annee_scolaire": self.annee.id, "classe": self.classe.id, "matiere": self.maths.id, "trimestre": 1,
                 "libelle": "Rattrapage", "bareme": "20", "poids": "1"}
        self.assertEqual(self.client.post(EVALS, corps).status_code, 400)
        self.assertEqual(self.client.delete(f"{EVALS}{self.evals['MATH'][0].id}/").status_code, 400)

    def test_reouverture_reservee_a_la_direction_avec_motif(self):
        self.valider(accepter_incomplet=True)
        corps = {"classe": self.classe.id, "periode": 1}
        pedago = client_de(utilisateur(Role.RESPONSABLE_PEDAGOGIQUE, "peda"))
        self.assertEqual(pedago.post(RESULTATS + "rouvrir/", {**corps, "motif": "x"}, format="json").status_code, 403)
        self.assertEqual(self.client.post(RESULTATS + "rouvrir/", corps, format="json").status_code, 400)
        self.assertEqual(self.client.post(RESULTATS + "rouvrir/", {**corps, "motif": "erreur"}, format="json").status_code, 200)
        self.assertFalse(self.resultats()["validee"])
        self.assertEqual(self.modifier_note(self.client).status_code, 200)

    def test_appreciations_et_decision(self):
        corps = {"classe": self.classe.id, "periode": 1, "appreciations": [
            {"etudiant": self.awa.id, "decision": "ADMIS", "observation_professeurs": "Peut mieux faire",
             "observation_directeur": "Encouragements"}]}
        r = self.client.post(RESULTATS + "appreciations/", corps, format="json")
        self.assertEqual(r.status_code, 200, r.content)
        awa = self.ligne(r.json(), self.awa)
        self.assertEqual((awa["decision"], awa["observation_directeur"]), ("ADMIS", "Encouragements"))
        corps["appreciations"][0]["decision"] = "PASSE"
        self.assertEqual(self.client.post(RESULTATS + "appreciations/", corps, format="json").status_code, 400)
        corps["appreciations"][0]["decision"] = "REDOUBLE"
        self.valider(accepter_incomplet=True)
        self.assertEqual(self.client.post(RESULTATS + "appreciations/", corps, format="json").status_code, 400)
        self.assertEqual(Appreciation.objects.get().decision, "ADMIS")


class BulletinsEtExportsTests(BaseEvaluations):
    def pdf(self, **params):
        with patch("reportlab.rl_config.pageCompression", 0):
            return self.client.get(RESULTATS + "bulletins/", {"classe": self.classe.id, "periode": 1, **params})

    def test_bulletins_provisoires_puis_valides(self):
        self.jeu_complet()
        provisoire = self.pdf()
        self.assertEqual(provisoire["Content-Type"], "application/pdf")
        self.assertTrue(provisoire.content.startswith(b"%PDF"))
        self.assertIn(b"PROVISOIRE", provisoire.content)
        self.client.post(RESULTATS + "valider/", {"classe": self.classe.id, "periode": 1, "accepter_incomplet": True}, format="json")
        self.assertNotIn(b"PROVISOIRE", self.pdf().content)

    def test_un_seul_eleve_et_bulletin_annuel(self):
        self.jeu_complet()
        un = self.pdf(etudiants=str(self.awa.id))
        self.assertEqual(un.status_code, 200)
        self.assertLess(len(un.content), len(self.pdf().content))
        self.assertEqual(self.client.get(RESULTATS + "bulletins/", {"classe": self.classe.id, "periode": 0}).status_code, 200)
        self.assertEqual(self.pdf(etudiants="abc").status_code, 400)

    def test_noms_avec_caracteres_speciaux_ne_cassent_pas_le_pdf(self):
        self.jeu_complet()
        sp = self.eleve("E7", "A&B <i>Test</i>", "Zoe")
        evaluation = Evaluation.objects.get(matiere=self.maths, libelle="Composition 1")
        self.noter(evaluation, sp, 12)
        self.assertEqual(self.pdf().status_code, 200)

    def test_export_du_classement(self):
        self.jeu_complet()
        r = self.client.get(RESULTATS + "exporter/", {"classe": self.classe.id, "periode": 1})
        feuille = openpyxl.load_workbook(io.BytesIO(r.content)).active
        valeurs = [[c.value for c in ligne] for ligne in feuille.iter_rows()]
        entete = next(l for l in valeurs if l[0] == "Rang")
        premier = valeurs[valeurs.index(entete) + 1]
        self.assertEqual((premier[0], premier[2]), (1, "Ba Moussa"))
        self.assertIn("Moyenne generale", entete)
        self.assertEqual(premier[entete.index("Moyenne generale")], 13.77)


class ImportNotesTests(BaseEvaluations):
    def setUp(self):
        super().setUp()
        self.evaluation_ = self.evaluation(self.maths, "Devoir", 1)
        self.base = f"{EVALS}{self.evaluation_.id}/"

    def modele(self):
        r = self.client.get(self.base + "modele-import/")
        self.assertEqual(r.status_code, 200)
        return openpyxl.load_workbook(io.BytesIO(r.content))

    def remplir(self, classeur, valeurs):
        for ligne in classeur.active.iter_rows(min_row=9):
            if ligne[1].value in valeurs:
                ligne[3].value = valeurs[ligne[1].value]
        return classeur

    def envoyer(self, classeur, **extra):
        tampon = io.BytesIO()
        classeur.save(tampon)
        fichier = SimpleUploadedFile("notes.xlsx", tampon.getvalue())
        return self.client.post(self.base + "importer/", {"fichier": fichier, **extra}, format="multipart")

    def test_apercu_puis_confirmation(self):
        classeur = self.remplir(self.modele(), {"E1": 14.5, "E2": "12,25", "E3": "AJ"})
        apercu = self.envoyer(classeur)
        self.assertEqual(apercu.status_code, 200, apercu.content)
        self.assertEqual(apercu.json()["resume"]["nouveaux"], 3)
        self.assertEqual(Note.objects.count(), 0)
        r = self.envoyer(classeur, confirmer="true")
        self.assertEqual(r.status_code, 201, r.content)
        notes = {n.etudiant.matricule: (n.statut, n.valeur) for n in Note.objects.all()}
        self.assertEqual(notes, {"E1": ("SAISIE", D("14.50")), "E2": ("SAISIE", D("12.25")), "E3": ("ABSENT_JUSTIFIE", None)})
        self.assertTrue(JournalAudit.objects.filter(action="IMPORT_NOTES").exists())

    def test_anomalies_bloquent_tout_et_cellule_vide_ne_change_rien(self):
        classeur = self.remplir(self.modele(), {"E1": 25, "E2": "abc"})       # hors bareme / illisible ; E3 vide
        classeur.active.append([None, "XXXX", "Inconnu", 10])
        r = self.envoyer(classeur, confirmer="true")
        self.assertEqual(r.status_code, 400)
        anomalies = {l["matricule"]: " ".join(l["anomalies"]) for l in r.json()["lignes"] if l["anomalies"]}
        self.assertIn("hors du bareme", anomalies["E1"])
        self.assertIn("note invalide", anomalies["E2"])
        self.assertIn("inconnu", anomalies["XXXX"])
        self.assertEqual(Note.objects.count(), 0)

    def test_aucun_ecrasement_silencieux(self):
        self.noter(self.evaluation_, self.awa, 10)
        classeur = self.remplir(self.modele(), {"E1": 15, "E2": 12})
        self.assertEqual(self.envoyer(classeur).json()["resume"]["modifies"], 1)
        self.assertEqual(self.envoyer(classeur, confirmer="true").status_code, 400)
        self.assertEqual(Note.objects.get(etudiant=self.awa).valeur, D("10"))
        self.assertEqual(self.envoyer(classeur, confirmer="true", remplacer="true").status_code, 201)
        self.assertEqual(Note.objects.get(etudiant=self.awa).valeur, D("15"))
        self.assertTrue(JournalAudit.objects.filter(action="MODIFICATION_NOTE").exists())

    def test_fichier_d_une_autre_evaluation_refuse(self):
        classeur = self.remplir(self.modele(), {"E1": 15})
        classeur.active["B5"].value = "Autre devoir"
        r = self.envoyer(classeur, confirmer="true")
        self.assertEqual(r.status_code, 400)
        self.assertIn("evaluation du fichier", r.json()["detail"])

    def test_export_reimportable(self):
        self.noter(self.evaluation_, self.awa, 13.5)
        feuille = openpyxl.load_workbook(io.BytesIO(self.client.get(self.base + "exporter/").content)).active
        ligne = next(l for l in feuille.iter_rows(min_row=9) if l[1].value == "E1")
        self.assertEqual(ligne[3].value, 13.5)


class PermissionsEvaluationsTests(BaseEvaluations):
    def setUp(self):
        super().setUp()
        self.prof = Professeur.objects.create(matricule="P1", nom="Sy", prenom="Oumar", tarif_horaire=100)
        AffectationProfesseur.objects.create(professeur=self.prof, matiere=self.maths, classe=self.classe, annee_scolaire=self.annee)
        self.client_prof = client_de(utilisateur(Role.PROFESSEUR, "prof", professeur=self.prof))

    def test_professeur_ne_voit_que_ses_evaluations(self):
        mienne = self.evaluation(self.maths, "Devoir", 1)
        self.evaluation(self.fr, "Devoir", 1)
        ids = [e["id"] for e in self.client_prof.get(EVALS).json()["resultats"]]
        self.assertEqual(ids, [mienne.id])

    def test_professeur_cree_seulement_pour_ses_affectations_et_saisit_ses_notes(self):
        corps = {"annee_scolaire": self.annee.id, "classe": self.classe.id, "matiere": self.maths.id,
                 "trimestre": 1, "libelle": "Devoir 1", "bareme": "20", "poids": "1"}
        r = self.client_prof.post(EVALS, corps)
        self.assertEqual(r.status_code, 201, r.content)
        self.assertEqual(self.client_prof.post(EVALS, {**corps, "matiere": self.fr.id}).status_code, 400)
        notes = {"notes": [{"etudiant": self.awa.id, "statut": "SAISIE", "valeur": "15"}]}
        self.assertEqual(self.client_prof.post(f"{EVALS}{r.json()['id']}/notes/", notes, format="json").status_code, 200)
        autre = self.evaluation(self.fr, "Devoir", 1)
        self.assertEqual(self.client_prof.post(f"{EVALS}{autre.id}/notes/", notes, format="json").status_code, 404)

    def test_roles_sans_acces(self):
        self.assertEqual(client_de(utilisateur(Role.COMPTABLE, "compta")).get(EVALS).status_code, 403)
        scolarite = client_de(utilisateur(Role.RESPONSABLE_SCOLARITE, "scol"))
        self.assertEqual(scolarite.get(EVALS).status_code, 403)                       # pas de saisie de notes
        self.assertEqual(scolarite.get(RESULTATS, {"classe": self.classe.id}).status_code, 200)   # mais lecture des resultats
        self.assertEqual(scolarite.post(RESULTATS + "valider/", {"classe": self.classe.id, "periode": 1}, format="json").status_code, 403)
        self.assertEqual(self.client_prof.get(RESULTATS, {"classe": self.classe.id}).status_code, 403)
