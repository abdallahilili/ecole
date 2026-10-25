"""Bilinguisme francais / arabe : traduction des messages, documents PDF/Excel, WhatsApp, noms arabes."""

import ast
import io
import os
from datetime import date, time

import openpyxl
from django.test import SimpleTestCase, TestCase
from rest_framework.test import APIClient

from apps.commun import pdf as pdf_commun
from apps.commun.constantes import Role
from apps.commun.langue import en_langue, est_arabe
from apps.commun.traduction import traduire
from apps.etudiants.models import Etudiant, Inscription
from apps.evaluations import bulletins, resultats
from apps.evaluations.models import Evaluation, Note
from apps.presences import exports as exports_presences
from apps.presences import services as services_presences
from apps.presences.models import Seance
from apps.professeurs.models import Professeur
from apps.referentiels.models import AffectationMatiereClasse, AnneeScolaire, Classe, Matiere, Niveau
from apps.utilisateurs.models import Utilisateur

RACINE_APPS = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def _gabarit(noeud):
    """Texte d'un litteral, d'une f-string (valeurs -> {}) ou d'une concatenation de litteraux ; None sinon."""
    if isinstance(noeud, ast.Constant) and isinstance(noeud.value, str):
        return noeud.value
    if isinstance(noeud, ast.JoinedStr):
        return "".join(v.value if isinstance(v, ast.Constant) else "{}" for v in noeud.values)
    if isinstance(noeud, ast.BinOp) and isinstance(noeud.op, ast.Add):
        gauche, droite = _gabarit(noeud.left), _gabarit(noeud.right)
        return gauche + droite if gauche is not None and droite is not None else None
    return None


def messages_du_code():
    """Messages d'erreur ecrits en dur : arguments de ErreurMetier(...) et de ValidationError(...) dans le code metier."""
    trouves = {}
    for dossier, _, fichiers in os.walk(RACINE_APPS):
        if "tests" in dossier or "migrations" in dossier:
            continue
        for fichier in fichiers:
            if not fichier.endswith(".py"):
                continue
            chemin = os.path.join(dossier, fichier)
            arbre = ast.parse(open(chemin, encoding="utf-8-sig").read())
            for noeud in ast.walk(arbre):
                if isinstance(noeud, ast.Call) and noeud.args:
                    nom = noeud.func.id if isinstance(noeud.func, ast.Name) else getattr(noeud.func, "attr", "")
                    if nom in ("ErreurMetier", "ValidationError"):
                        texte = _gabarit(noeud.args[0])
                        if texte and any(c.isalpha() for c in texte.replace("{}", "")):
                            trouves[texte] = f"{os.path.relpath(chemin, RACINE_APPS)}:{noeud.lineno}"
    return trouves


class MoteurTraductionTests(SimpleTestCase):
    def test_le_francais_est_renvoye_tel_quel(self):
        self.assertEqual(traduire("Resultats valides."), "Resultats valides.")
        with en_langue("fr"):
            self.assertEqual(traduire("Resultats valides."), "Resultats valides.")

    def test_gabarit_avec_valeurs_et_jours(self):
        with en_langue("ar"):
            texte = traduire("La classe 6AF a deux cours en meme temps (mardi 09:00-10:00) : H/G - 6AF et Francais - 6AF.")
        self.assertIn("6AF", texte)
        self.assertIn("الثلاثاء", texte)
        self.assertNotIn("a deux cours", texte)

    def test_un_gabarit_n_avale_pas_plusieurs_messages(self):
        with en_langue("ar"):
            texte = traduire("Presences refusees : statut invalide 'X' pour l'etudiant 12 ; l'etudiant 5 est present deux fois")
        self.assertNotIn("est present deux fois", texte)
        self.assertNotIn("statut invalide", texte)

    def test_texte_inconnu_inchange(self):
        with en_langue("ar"):
            self.assertEqual(traduire("Une phrase que personne n'a traduite"), "Une phrase que personne n'a traduite")

    def test_contexte_de_langue_restaure(self):
        self.assertFalse(est_arabe())
        with en_langue("ar"):
            self.assertTrue(est_arabe())
        self.assertFalse(est_arabe())

    def test_tous_les_messages_d_erreur_du_code_ont_une_traduction(self):
        """Tout message ecrit en dur dans ErreurMetier / ValidationError doit avoir sa traduction arabe."""
        sans_traduction = []
        with en_langue("ar"):
            for gabarit, origine in messages_du_code().items():
                exemple = gabarit.replace("{}", "XX")
                if traduire(exemple) == exemple:
                    sans_traduction.append(f"{origine} : {gabarit}")
        self.assertEqual(sans_traduction, [], "Messages sans traduction arabe :\n" + "\n".join(sans_traduction))


class ArabeVisuelTests(SimpleTestCase):
    def test_suite_numerique_garde_son_ordre(self):
        with en_langue("ar"):
            visuel = pdf_commun._visuel("السنة الدراسية : 2026-2027")
        self.assertIn("2026-2027", visuel)

    def test_texte_latin_inchange(self):
        self.assertEqual(pdf_commun._visuel("Abdallahi 2026"), "Abdallahi 2026")


class DonneesBilingues(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.annee = AnneeScolaire.objects.create(libelle="2026-2027", date_debut=date(2026, 9, 1), date_fin=date(2027, 6, 30), active=True)
        cls.niveau = Niveau.objects.create(nom="6eme", nom_ar="السنة السادسة", code="6E", ordre=1)
        cls.classe = Classe.objects.create(nom="6AF", niveau=cls.niveau, annee_scolaire=cls.annee, capacite=30)
        cls.maths = Matiere.objects.create(nom="Maths", nom_ar="الرياضيات", code="MAT")
        cls.sans_ar = Matiere.objects.create(nom="EPS", code="EPS")
        for matiere in (cls.maths, cls.sans_ar):
            AffectationMatiereClasse.objects.create(matiere=matiere, classe=cls.classe, coefficient=2)
        cls.etudiant = Etudiant.objects.create(matricule="E1", nom="Abdallahi", prenom="Aziza", telephone_pere="22334455")
        cls.etudiant_ar = Etudiant.objects.create(matricule="E2", nom="محمد", prenom="سالم")
        for e in (cls.etudiant, cls.etudiant_ar):
            Inscription.objects.create(etudiant=e, annee_scolaire=cls.annee, classe=cls.classe, date_inscription=date(2026, 9, 1))
        cls.prof = Professeur.objects.create(matricule="P1", nom="Sy", prenom="Oumar", tarif_horaire=900)
        cls.admin = Utilisateur.objects.create_user(identifiant="admin", mot_de_passe="Secret-2026-ab", nom="Admin", role=Role.ADMINISTRATEUR)
        evaluation = Evaluation.objects.create(annee_scolaire=cls.annee, classe=cls.classe, matiere=cls.maths, trimestre=1, libelle="Devoir", poids=1)
        Note.objects.create(evaluation=evaluation, etudiant=cls.etudiant, valeur=15)
        Note.objects.create(evaluation=evaluation, etudiant=cls.etudiant_ar, valeur=12)
        cls.seance = Seance.objects.create(
            annee_scolaire=cls.annee, classe=cls.classe, matiere=cls.maths, professeur=cls.prof, date=date(2026, 10, 5),
            heure_debut=time(8, 0), heure_fin=time(9, 0),
        )

    def api(self, langue):
        c = APIClient()
        c.force_authenticate(self.admin)
        c.defaults["HTTP_ACCEPT_LANGUAGE"] = langue
        return c


class MessagesApiTests(DonneesBilingues):
    def test_erreur_metier_traduite_selon_l_en_tete(self):
        url = "/api/referentiels/annees-scolaires/"
        corps = {"libelle": "2030-2031", "date_debut": "2030-09-01", "date_fin": "2030-08-01"}
        francais = self.api("fr").post(url, corps).json()
        arabe = self.api("ar").post(url, corps)
        self.assertEqual(arabe.status_code, 400)
        self.assertEqual(arabe["Content-Language"], "ar")
        self.assertNotEqual(francais, arabe.json())
        self.assertTrue(any("؀" <= c <= "ۿ" for c in str(arabe.json())), arabe.json())

    def test_defaut_francais_et_langue_inconnue(self):
        reponse = APIClient()
        reponse.force_authenticate(self.admin)
        r = reponse.post("/api/referentiels/annees-scolaires/", {"libelle": "x", "date_debut": "2030-09-01", "date_fin": "2030-08-01"})
        self.assertEqual(r["Content-Language"], "fr")
        r = self.api("de").post("/api/referentiels/annees-scolaires/", {"libelle": "x", "date_debut": "2030-09-01", "date_fin": "2030-08-01"})
        self.assertEqual(r["Content-Language"], "fr")

    def test_les_donnees_saisies_ne_sont_jamais_traduites(self):
        """Un etudiant qui s'appelle « Present » reste « Present » : seules les cles de message sont traduites."""
        Etudiant.objects.create(matricule="E3", nom="Present", prenom="Absent")
        liste = self.api("ar").get("/api/etudiants/etudiants/", {"recherche": "Present"}).json()["resultats"]
        self.assertEqual(liste[0]["nom"], "Present")

    def test_nom_arabe_des_matieres(self):
        arabe = self.api("ar").get("/api/referentiels/matieres/").json()["resultats"]
        francais = self.api("fr").get("/api/referentiels/matieres/").json()["resultats"]
        par_nom = lambda liste, nom: next(m for m in liste if m["nom"] == nom)
        self.assertEqual(par_nom(arabe, "Maths")["nom_affiche"], "الرياضيات")
        self.assertEqual(par_nom(arabe, "EPS")["nom_affiche"], "EPS")          # pas de nom arabe : repli sur le nom principal
        self.assertEqual(par_nom(francais, "Maths")["nom_affiche"], "Maths")
        self.assertEqual(par_nom(arabe, "Maths")["nom_ar"], "الرياضيات")


class DocumentsArabesTests(DonneesBilingues):
    def _pdf(self, url, langue="ar"):
        reponse = self.api(langue).get(url)
        self.assertEqual(reponse.status_code, 200, reponse.content[:200])
        self.assertEqual(reponse.content[:4], b"%PDF")
        return reponse.content

    def test_bulletins_en_arabe_et_en_francais(self):
        url = f"/api/evaluations/resultats/bulletins/?classe={self.classe.id}&periode=1"
        arabe, francais = self._pdf(url), self._pdf(url, "fr")
        self.assertNotEqual(arabe, francais)
        self.assertGreater(len(arabe), 5000)

    def test_bulletin_avec_noms_arabes_dans_un_document_francais(self):
        """Un eleve saisi en caracteres arabes s'imprime meme dans un document francais (police arabe de repli)."""
        donnees = resultats.resultats(self.classe, 1)
        octets = bulletins.pdf_bulletins(donnees, {self.etudiant_ar.id})
        self.assertEqual(octets[:4], b"%PDF")

    def test_liste_de_presence_emploi_et_fiche(self):
        self._pdf(f"/api/presences/seances/{self.seance.id}/pdf/")
        self._pdf("/api/presences/seances/modele-vierge/pdf/")

    def test_excel_de_droite_a_gauche_en_arabe(self):
        reponse = self.api("ar").get(f"/api/presences/seances/{self.seance.id}/excel/")
        classeur = openpyxl.load_workbook(io.BytesIO(reponse.content))
        feuille = classeur.active
        self.assertTrue(feuille.sheet_view.rightToLeft)
        textes = [str(c.value) for ligne in feuille.iter_rows() for c in ligne if c.value]
        self.assertTrue(any("الرقم المدرسي" in t for t in textes), textes[:20])
        francais = openpyxl.load_workbook(io.BytesIO(self.api("fr").get(f"/api/presences/seances/{self.seance.id}/excel/").content)).active
        self.assertFalse(bool(francais.sheet_view.rightToLeft))

    def test_modele_import_etudiants_en_arabe(self):
        reponse = self.api("ar").get("/api/etudiants/etudiants/modele-import/")
        feuille = openpyxl.load_workbook(io.BytesIO(reponse.content)).active
        self.assertEqual(feuille.cell(row=1, column=1).value, "الرقم المدرسي")


class WhatsAppBilingueTests(DonneesBilingues):
    def test_message_selon_la_langue_des_parents(self):
        from apps.commun.models import Parametre
        presence = self.seance.presences.create(etudiant=self.etudiant, statut="ABSENT") if hasattr(self.seance, "presences") else None
        self.assertIsNotNone(presence)
        self.assertIn("Bonjour", services_presences.texte_message(presence))
        Parametre.objects.create(cle="langue_whatsapp", valeur="ar")
        arabe = services_presences.texte_message(presence)
        self.assertIn("مرحبًا", arabe)
        self.assertIn("الرياضيات", arabe)       # matiere dans la langue du message
        self.assertIn("Abdallahi", arabe)


class ExportsListesArabesTests(DonneesBilingues):
    COLONNES = '[{"cle":"matricule","titre":"الرقم المدرسي"},{"cle":"nom","titre":"الاسم"},{"cle":"actif","titre":"نشط"}]'

    def _export(self, fichier):
        reponse = self.api("ar").get("/api/etudiants/etudiants/export/", {"fichier": fichier, "titre": "التلاميذ", "colonnes": self.COLONNES})
        self.assertEqual(reponse.status_code, 200, reponse.content[:200])
        return reponse

    def test_pdf_xlsx_csv_avec_titres_et_noms_arabes(self):
        self.assertEqual(self._export("pdf").content[:4], b"%PDF")
        classeur = openpyxl.load_workbook(io.BytesIO(self._export("xlsx").content))
        feuille = classeur.active
        self.assertTrue(feuille.sheet_view.rightToLeft)
        valeurs = [c.value for ligne in feuille.iter_rows() for c in ligne]
        self.assertIn("محمد", valeurs)                 # nom saisi en arabe conserve tel quel
        self.assertIn("نعم", valeurs)                  # « Oui » traduit
        csv = self._export("csv").content.decode("utf-8-sig")
        self.assertIn("الرقم المدرسي", csv)
        self.assertIn("محمد", csv)

    def test_tableau_de_bord_et_messages_d_importation_en_arabe(self):
        self.assertEqual(self.api("ar").get("/api/tableau-de-bord/").status_code, 200)
        reponse = self.api("ar").post("/api/etudiants/etudiants/importer/", {}, format="multipart")
        self.assertEqual(reponse.status_code, 400)
        self.assertIn("ملف", reponse.json()["detail"])


class CodesTraduitsDansLesExportsTests(DonneesBilingues):
    def test_sexe_et_statuts_lisibles_en_arabe(self):
        Etudiant.objects.filter(pk=self.etudiant.pk).update(sexe="F")
        cols = '[{"cle":"nom","titre":"الاسم"},{"cle":"sexe","titre":"الجنس"}]'
        r = self.api("ar").get("/api/etudiants/etudiants/export/", {"fichier": "csv", "colonnes": cols, "recherche": "Abdallahi"})
        texte = r.content.decode("utf-8-sig")
        self.assertIn("أنثى", texte)
        self.assertNotIn(";F", texte)
        r = self.api("fr").get("/api/etudiants/etudiants/export/", {"fichier": "csv", "colonnes": cols, "recherche": "Abdallahi"})
        self.assertIn("Feminin", r.content.decode("utf-8-sig"))
