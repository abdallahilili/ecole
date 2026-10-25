"""Listes de presence imprimables (PDF), Excel et modele d'import."""

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from reportlab.lib.units import cm
from reportlab.platypus import Spacer

from apps.commun.excel import orienter_feuille
from apps.commun.parametres import lire_parametre
from apps.commun.traduction import traduire
from apps.commun.pdf import (Tableau as Table, 
    STYLE_NORMAL, STYLE_TITRE, construire_pdf, entete_etablissement, p, style_tableau,
)
from .models import StatutPresence
from apps.etudiants.services import eleves_classe
from .services import liste_appel

LIGNES_MODELE_VIERGE = 40


def _nom_etablissement():
    from apps.commun.pdf import _parametre_localise
    return _parametre_localise("nom_etablissement")

# Disposition FIXE du modele d'import : l'importeur relit ces cellules pour
# verifier que le fichier correspond bien au cours choisi.
META_ANNEE, META_CLASSE, META_MATIERE, META_DATE, META_DEBUT, META_FIN = 1, 2, 3, 4, 5, 6
LIGNE_ENTETE_IMPORT = 8
PREMIERE_LIGNE_IMPORT = 9
CODES_STATUT = {"P": StatutPresence.PRESENT, "A": StatutPresence.ABSENT, "R": StatutPresence.RETARD}


def contexte_seance(seance):
    return {
        "annee": seance.annee_scolaire.libelle,
        "classe": seance.classe.nom,
        "matiere": seance.matiere.nom_affiche,
        "professeur": seance.professeur.nom_complet if seance.professeur else "",
        "date": seance.date.strftime("%d/%m/%Y"),
        "horaire": f"{seance.heure_debut:%H:%M} - {seance.heure_fin:%H:%M}",
        "lignes": liste_appel(seance),
    }


def contexte_vierge():
    return {
        "annee": "", "classe": "", "matiere": "", "professeur": "", "date": "", "horaire": "",
        "lignes": [
            {"numero": i, "matricule": "", "nom_complet": "", "statut": None, "observation": ""}
            for i in range(1, LIGNES_MODELE_VIERGE + 1)
        ],
    }


def _marque(ligne, statut):
    return "X" if ligne["statut"] == statut else ""


def pdf_liste(contexte):
    infos = Table(
        [
            [p(f"Annee scolaire : {contexte['annee']}"), p(f"Classe : {contexte['classe']}")],
            [p(f"Matiere : {contexte['matiere']}"), p(f"Professeur : {contexte['professeur']}")],
            [p(f"Date : {contexte['date']}"), p(f"Horaire : {contexte['horaire']}")],
        ],
        colWidths=[9 * cm, 9 * cm],
    )
    donnees = [["N°", "Matricule", "Nom complet", "Present", "Absent", "Retard", "Observations"]]
    for ligne in contexte["lignes"]:
        donnees.append([
            ligne["numero"], p(ligne["matricule"]), p(ligne["nom_complet"]),
            _marque(ligne, StatutPresence.PRESENT), _marque(ligne, StatutPresence.ABSENT),
            _marque(ligne, StatutPresence.RETARD), p(ligne["observation"]),
        ])
    tableau = Table(
        donnees, repeatRows=1, rowHeights=[0.7 * cm] + [0.75 * cm] * len(contexte["lignes"]),
        colWidths=[0.9 * cm, 2.4 * cm, 5.6 * cm, 1.5 * cm, 1.5 * cm, 1.5 * cm, 4.6 * cm],
    )
    tableau.setStyle(style_tableau())
    elements = entete_etablissement() + [
        p("LISTE DE PRESENCE", STYLE_TITRE), infos, Spacer(1, 8), tableau, Spacer(1, 14),
        p("Signature du professeur : ______________________", STYLE_NORMAL),
    ]
    return construire_pdf(elements, titre="Liste de presence")


def xlsx_liste(contexte):
    classeur = openpyxl.Workbook()
    feuille = classeur.active
    feuille.title = traduire("Liste de presence")
    orienter_feuille(feuille)
    feuille.append([_nom_etablissement()])
    feuille.append([traduire("LISTE DE PRESENCE")])
    feuille["A1"].font = Font(bold=True, size=13)
    feuille["A2"].font = Font(bold=True, size=12)
    for etiquette, valeur in (
        ("Annee scolaire", contexte["annee"]), ("Classe", contexte["classe"]),
        ("Matiere", contexte["matiere"]), ("Professeur", contexte["professeur"]),
        ("Date", contexte["date"]), ("Horaire", contexte["horaire"]),
    ):
        feuille.append([traduire(etiquette), valeur])
    feuille.append([])
    entete = [traduire(h) for h in ["N°", "Matricule", "Nom complet", "Present", "Absent", "Retard", "Observations"]]
    feuille.append(entete)
    ligne_entete = feuille.max_row
    for index in range(1, len(entete) + 1):
        cellule = feuille.cell(row=ligne_entete, column=index)
        cellule.font = Font(bold=True)
        cellule.fill = PatternFill("solid", fgColor="EEF1F3")
    for ligne in contexte["lignes"]:
        feuille.append([
            ligne["numero"], ligne["matricule"], ligne["nom_complet"],
            _marque(ligne, StatutPresence.PRESENT), _marque(ligne, StatutPresence.ABSENT),
            _marque(ligne, StatutPresence.RETARD), ligne["observation"],
        ])
    for index, largeur in enumerate([6, 14, 34, 10, 10, 10, 36], start=1):
        feuille.column_dimensions[get_column_letter(index)].width = largeur
    return classeur


# --------------------------------------------------- modele de presence d'une classe

JOURS_MODELE = 31


def contexte_classe(classe):
    """Eleves (inscriptions actives) de la classe pour le modele de presence mensuel."""
    eleves = eleves_classe(classe.pk, classe.annee_scolaire_id)
    return {
        "annee": classe.annee_scolaire.libelle,
        "classe": classe.nom,
        "lignes": [
            {"numero": i, "matricule": inscription.etudiant.matricule, "nom_complet": inscription.etudiant.nom_complet}
            for i, inscription in enumerate(eleves, start=1)
        ],
    }


def pdf_modele_classe(contexte):
    """Feuille de presence mensuelle imprimable (paysage) : une ligne par eleve, une colonne par jour du mois,
    totaux d'absences et de retards, et les mentions a remplir a la main (mois, matiere, professeur)."""
    largeur = 27.25 * cm
    infos = Table(
        [
            [p(f"Annee scolaire : {contexte['annee']}"), p(f"Classe : {contexte['classe']}"),
             p(f"Effectif : {len(contexte['lignes'])}")],
            [p("Mois : ____________________"), p("Matiere : ____________________"), p("Professeur : ____________________")],
        ],
        colWidths=[largeur / 3] * 3,
    )
    jours = [str(j) for j in range(1, JOURS_MODELE + 1)]
    donnees = [["N°", "Matricule", "Nom complet"] + jours + ["Abs.", "Ret."]]
    for ligne in contexte["lignes"]:
        donnees.append([ligne["numero"], p(ligne["matricule"]), p(ligne["nom_complet"])] + [""] * JOURS_MODELE + ["", ""])
    largeurs = [0.8 * cm, 2.1 * cm, 5.3 * cm] + [0.55 * cm] * JOURS_MODELE + [1 * cm, 1 * cm]
    tableau = Table(donnees, repeatRows=1, colWidths=largeurs, rowHeights=[0.7 * cm] + [0.62 * cm] * len(contexte["lignes"]))
    style = style_tableau()
    premiere_jour, derniere_jour = 3, 2 + JOURS_MODELE
    style.add("FONTSIZE", (premiere_jour, 0), (derniere_jour, 0), 6.5)
    style.add("ALIGN", (premiere_jour, 0), (derniere_jour, 0), "CENTER")
    style.add("LEFTPADDING", (premiere_jour, 0), (derniere_jour, -1), 0)
    style.add("RIGHTPADDING", (premiere_jour, 0), (derniere_jour, -1), 0)
    tableau.setStyle(style)
    elements = entete_etablissement(sum(largeurs)) + [
        p("REGISTRE DE PRESENCE MENSUEL", STYLE_TITRE), infos, Spacer(1, 6), tableau, Spacer(1, 8),
        p("Legende : P = present, A = absent, R = retard", STYLE_NORMAL), Spacer(1, 8),
        p("Signature du professeur : ______________________          Visa du directeur : ______________________", STYLE_NORMAL),
    ]
    return construire_pdf(elements, paysage=True, titre="Registre de presence")


def xlsx_modele_classe(contexte):
    classeur = openpyxl.Workbook()
    feuille = classeur.active
    feuille.title = traduire("Registre de presence")[:31]
    orienter_feuille(feuille)
    feuille.append([_nom_etablissement()])
    feuille.append([traduire("REGISTRE DE PRESENCE MENSUEL")])
    feuille["A1"].font = Font(bold=True, size=13)
    feuille["A2"].font = Font(bold=True, size=12)
    for etiquette, valeur in (
        ("Annee scolaire", contexte["annee"]), ("Classe", contexte["classe"]), ("Effectif", len(contexte["lignes"])),
        ("Mois", ""), ("Matiere", ""), ("Professeur", ""),
    ):
        feuille.append([traduire(etiquette), valeur])
    feuille.append([])
    entete = [traduire(h) for h in ["N°", "Matricule", "Nom complet"]] + list(range(1, JOURS_MODELE + 1)) + [traduire("Abs."), traduire("Ret.")]
    feuille.append(entete)
    ligne_entete = feuille.max_row
    for index in range(1, len(entete) + 1):
        cellule = feuille.cell(row=ligne_entete, column=index)
        cellule.font = Font(bold=True)
        cellule.fill = PatternFill("solid", fgColor="EEF1F3")
        cellule.alignment = Alignment(horizontal="center")
    for ligne in contexte["lignes"]:
        feuille.append([ligne["numero"], ligne["matricule"], ligne["nom_complet"]])
    feuille.append([])
    feuille.append([traduire("Legende : P = present, A = absent, R = retard")])
    for index, largeur in enumerate([6, 14, 34] + [4.5] * JOURS_MODELE + [7, 7], start=1):
        feuille.column_dimensions[get_column_letter(index)].width = largeur
    feuille.freeze_panes = feuille.cell(row=ligne_entete + 1, column=4)
    return classeur


def xlsx_modele_import(seance_ctx):
    """Modele Excel a remplir : cours en en-tete, une ligne par eleve de la classe."""
    classeur = openpyxl.Workbook()
    feuille = classeur.active
    feuille.title = traduire("Presences")
    orienter_feuille(feuille)
    meta = [
        ("Annee scolaire", seance_ctx["annee"]), ("Classe", seance_ctx["classe"]),
        ("Matiere (code)", seance_ctx["matiere_code"]), ("Date (AAAA-MM-JJ)", seance_ctx["date_iso"]),
        ("Heure debut (HH:MM)", seance_ctx["heure_debut"]), ("Heure fin (HH:MM)", seance_ctx["heure_fin"]),
    ]
    for index, (etiquette, valeur) in enumerate(meta, start=1):
        feuille.cell(row=index, column=1, value=traduire(etiquette)).font = Font(bold=True)
        feuille.cell(row=index, column=2, value=valeur)
    feuille.cell(row=7, column=1, value=traduire("Statut : P = present, A = absent, R = retard. Ne pas modifier les lignes 1 a 6."))
    feuille.cell(row=7, column=1).font = Font(italic=True)
    entete = [traduire(h) for h in ["N°", "Matricule", "Nom complet", "Statut (P/A/R)", "Observation"]]
    for index, titre in enumerate(entete, start=1):
        cellule = feuille.cell(row=LIGNE_ENTETE_IMPORT, column=index, value=titre)
        cellule.font = Font(bold=True)
        cellule.fill = PatternFill("solid", fgColor="EEF1F3")
    for decalage, ligne in enumerate(seance_ctx["lignes"]):
        rang = PREMIERE_LIGNE_IMPORT + decalage
        feuille.cell(row=rang, column=1, value=ligne["numero"])
        feuille.cell(row=rang, column=2, value=ligne["matricule"])
        feuille.cell(row=rang, column=3, value=ligne["nom_complet"])
        statut = next((code for code, nom in CODES_STATUT.items() if nom == ligne["statut"]), None)
        feuille.cell(row=rang, column=4, value=statut)
        feuille.cell(row=rang, column=5, value=ligne["observation"] or None)
    derniere = PREMIERE_LIGNE_IMPORT + max(len(seance_ctx["lignes"]), 1) - 1
    controle = DataValidation(type="list", formula1='"P,A,R"', allow_blank=True)
    feuille.add_data_validation(controle)
    controle.add(f"D{PREMIERE_LIGNE_IMPORT}:D{derniere}")
    for index, largeur in enumerate([22, 16, 34, 16, 36], start=1):
        feuille.column_dimensions[get_column_letter(index)].width = largeur
    for ligne in feuille.iter_rows(min_row=PREMIERE_LIGNE_IMPORT, max_row=derniere, min_col=4, max_col=4):
        for cellule in ligne:
            cellule.alignment = Alignment(horizontal="center")
    return classeur
